"""Generate the example workbooks and their xlsx2txt exports.

Run from the repository root:

    python examples/generate.py

Every example is built with openpyxl and saved with fixed document dates, so
running the script again produces the same content.
"""

import datetime
import io
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.packaging.custom import BoolProperty, StringProperty
from openpyxl.packaging.relationship import Relationship
from openpyxl.pivot.cache import CacheDefinition
from openpyxl.pivot.record import RecordList
from openpyxl.pivot.table import TableDefinition
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.external_link.external import (
    ExternalBook,
    ExternalCell,
    ExternalLink,
    ExternalRow,
    ExternalSheetData,
    ExternalSheetDataSet,
    ExternalSheetNames,
)
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.xml.functions import fromstring

EXAMPLES_DIR = Path(__file__).resolve().parent
FIXED_DATE = datetime.datetime(2026, 1, 1, 12, 0, 0)

PRODUCTS = [
    ("Widget", "Hardware", 120, 9.99, datetime.date(2026, 1, 5), True),
    ("Gadget", "Hardware", 45, 24.50, datetime.date(2026, 1, 12), True),
    ("Gizmo", "Hardware", 80, 14.25, datetime.date(2026, 2, 3), False),
    ("Support plan", "Services", 12, 199.00, datetime.date(2026, 2, 20), True),
    ("Training", "Services", 5, 450.00, datetime.date(2026, 3, 9), False),
    ("Консультація", "Послуги", 3, 1200.00, datetime.date(2026, 3, 30), True),
]
HEADER_FONT = Font(bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF305496")


def _new_workbook(title: str) -> Workbook:
    wb = Workbook()
    wb.properties.title = title
    wb.properties.creator = "xlsx2txt examples"
    wb.properties.created = FIXED_DATE
    return wb


def _header(ws, values) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _png(color, size) -> bytes:
    from PIL import Image as PILImage

    buffer = io.BytesIO()
    PILImage.new("RGB", size, color).save(buffer, format="PNG")
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Examples
# ---------------------------------------------------------------------------

def basic_data() -> Workbook:
    """Plain values of every type, column widths, freeze panes and a filter."""
    wb = _new_workbook("Basic data")
    ws = wb.active
    ws.title = "Sales"
    _header(ws, ["Product", "Category", "Quantity", "Price", "Date", "In stock"])
    for row in PRODUCTS:
        ws.append(row)
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "#,##0.00"
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = "yyyy-mm-dd"
    for column, width in zip("ABCDEF", (18, 12, 10, 10, 12, 10)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{ws.max_row}"
    return wb


def styles() -> Workbook:
    """Fonts, fills, borders, alignment, number formats, merged cells,
    conditional formatting and data validation."""
    wb = _new_workbook("Styles")
    ws = wb.active
    ws.title = "Styles"

    ws.merge_cells("A1:E1")
    ws["A1"] = "Quarterly report"
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color="FF1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    thin = Side(style="thin", color="FF999999")
    _header(ws, ["Quarter", "Revenue", "Costs", "Margin", "Status"])
    rows = [("Q1", 125000, 98000), ("Q2", 142000, 101000), ("Q3", 98000, 105000), ("Q4", 171000, 120000)]
    for quarter, revenue, costs in rows:
        ws.append([quarter, revenue, costs, (revenue - costs) / revenue, "open"])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=5):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        row[1].number_format = '#,##0 "UAH"'
        row[2].number_format = '#,##0 "UAH"'
        row[3].number_format = "0.0%"

    ws["A8"] = "Wrapped text that does not fit into one line of the cell"
    ws["A8"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B8"] = "Rotated"
    ws["B8"].alignment = Alignment(text_rotation=45)
    ws["C8"] = "Italic strike"
    ws["C8"].font = Font(italic=True, strike=True)
    ws["D8"] = "Theme color"
    ws["D8"].fill = PatternFill("solid", fgColor="FFFFF2CC")

    ws.conditional_formatting.add("D3:D6", CellIsRule(operator="lessThan", formula=["0"],
                                                       font=Font(color="FF9C0006"),
                                                       fill=PatternFill("solid", bgColor="FFFFC7CE")))
    ws.conditional_formatting.add("B3:B6", DataBarRule(start_type="min", end_type="max", color="FF638EC6"))
    ws.conditional_formatting.add("C3:C6", ColorScaleRule(start_type="min", start_color="FF63BE7B",
                                                          end_type="max", end_color="FFF8696B"))

    status = DataValidation(type="list", formula1='"open,closed,archived"', allow_blank=False,
                            showErrorMessage=True, error="Choose a status from the list")
    status.add("E3:E6")
    ws.add_data_validation(status)

    for column, width in zip("ABCDE", (14, 16, 16, 10, 12)):
        ws.column_dimensions[column].width = width
    return wb


def formulas() -> Workbook:
    """Formulas, an array formula, named ranges, a table, a hidden sheet,
    comments and hyperlinks."""
    wb = _new_workbook("Formulas")
    ws = wb.active
    ws.title = "Orders"
    _header(ws, ["Product", "Quantity", "Price", "Total"])
    for index, (name, _, quantity, price, _, _) in enumerate(PRODUCTS[:5], start=2):
        ws.append([name, quantity, price, f"=B{index}*C{index}"])
    last = ws.max_row
    ws[f"A{last + 1}"] = "Sum"
    ws[f"D{last + 1}"] = f"=SUM(D2:D{last})"
    ws[f"A{last + 2}"] = "Average price"
    ws[f"D{last + 2}"] = f"=AVERAGE(C2:C{last})"
    ws[f"A{last + 3}"] = "With VAT"
    ws[f"D{last + 3}"] = f"=D{last + 1}*(1+VAT)"
    ws["F1"] = "Doubled"
    ws["F2"] = ArrayFormula(f"F2:F{last}", f"=B2:B{last}*2")

    table = Table(displayName="Orders", ref=f"A1:D{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    ws["A1"].comment = Comment("Product names come from the catalog", "Analyst")
    ws[f"A{last + 5}"] = "Documentation"
    ws[f"A{last + 5}"].hyperlink = "https://github.com/click0/xlsx2txt"
    ws[f"A{last + 5}"].style = "Hyperlink"

    settings = wb.create_sheet("Settings")
    settings["A1"] = "VAT"
    settings["B1"] = 0.2
    settings.sheet_state = "hidden"
    wb.defined_names["VAT"] = DefinedName("VAT", attr_text="Settings!$B$1")

    summary = wb.create_sheet("Summary")
    summary["A1"] = "Orders total"
    summary["B1"] = f"=Orders!D{last + 1}"
    summary["A2"] = "Largest order"
    summary["B2"] = f"=MAX(Orders!D2:D{last})"
    return wb


def images_and_charts() -> Workbook:
    """An image, a bar and a line chart on a sheet, and a chart sheet."""
    wb = _new_workbook("Images and charts")
    ws = wb.active
    ws.title = "Data"
    _header(ws, ["Month", "Revenue", "Costs"])
    for month, revenue, costs in [("Jan", 42, 30), ("Feb", 47, 31), ("Mar", 51, 35), ("Apr", 49, 33)]:
        ws.append([month, revenue, costs])

    ws.add_image(Image(io.BytesIO(_png((48, 84, 150), (120, 40)))), "E1")

    categories = Reference(ws, min_col=1, min_row=2, max_row=5)
    bar = BarChart()
    bar.title = "Revenue and costs"
    bar.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=5), titles_from_data=True)
    bar.set_categories(categories)
    ws.add_chart(bar, "E4")

    line = LineChart()
    line.title = "Revenue trend"
    line.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    line.set_categories(categories)
    ws.add_chart(line, "E20")

    pie = PieChart()
    pie.title = "Revenue share"
    pie.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    pie.set_categories(categories)
    wb.create_chartsheet("Pie chart").add_chart(pie)
    return wb


def rich_text() -> Workbook:
    """Cells with several differently formatted runs of text."""
    wb = _new_workbook("Rich text")
    ws = wb.active
    ws.title = "Notes"
    ws["A1"] = CellRichText([
        "Status: ",
        TextBlock(InlineFont(b=True, color="FF00B050"), "approved"),
        " by ",
        TextBlock(InlineFont(i=True), "finance"),
    ])
    ws["A2"] = CellRichText([
        TextBlock(InlineFont(rFont="Courier New", sz=10), "E=mc"),
        TextBlock(InlineFont(rFont="Courier New", sz=10, vertAlign="superscript"), "2"),
    ])
    ws["A3"] = CellRichText([
        "Увага: ",
        TextBlock(InlineFont(b=True, u="single", color="FFC00000"), "термін до 31.03"),
    ])
    ws.column_dimensions["A"].width = 40
    return wb


def workbook_features() -> Workbook:
    """Link to another workbook, custom properties, protection and printing."""
    wb = _new_workbook("Workbook features")
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "Budget from another file"
    ws["B1"] = "=[1]Budget!B2"
    ws["A2"] = "Protected sheet, print settings: landscape, fit to width"

    book = ExternalBook(
        sheetNames=ExternalSheetNames(sheetName=["Budget"]),
        sheetDataSet=ExternalSheetDataSet(sheetData=[
            ExternalSheetData(sheetId=0, row=[ExternalRow(r=2, cell=[ExternalCell(r="B2", v="250000")])]),
        ]),
        id="rId1",
    )
    link = ExternalLink(externalBook=book)
    link.file_link = Relationship(type="externalLinkPath", Target="budget-2026.xlsx", TargetMode="External")
    wb._external_links.append(link)

    wb.custom_doc_props.append(StringProperty(name="Department", value="Finance"))
    wb.custom_doc_props.append(BoolProperty(name="Approved", value=True))

    ws.protection.sheet = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.sheet_properties.tabColor = "FFC00000"
    ws.column_dimensions["A"].width = 60
    return wb


# openpyxl cannot create pivot tables, so this one is described in XML the way
# Excel stores it: cache definition, cached records and the table itself.
PIVOT_MAIN = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
PIVOT_REL = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'

PIVOT_DATA = [("Category", "Amount"), ("Fruit", 1), ("Veg", 2), ("Fruit", 3), ("Veg", 4)]

PIVOT_CACHE = f"""<pivotCacheDefinition {PIVOT_MAIN} {PIVOT_REL} refreshOnLoad="1" recordCount="4">
  <cacheSource type="worksheet"><worksheetSource ref="A1:B5" sheet="Data"/></cacheSource>
  <cacheFields count="2">
    <cacheField name="Category" numFmtId="0">
      <sharedItems count="2"><s v="Fruit"/><s v="Veg"/></sharedItems>
    </cacheField>
    <cacheField name="Amount" numFmtId="0">
      <sharedItems containsSemiMixedTypes="0" containsString="0" containsNumber="1"
                   containsInteger="1" minValue="1" maxValue="4"/>
    </cacheField>
  </cacheFields>
</pivotCacheDefinition>"""

PIVOT_RECORDS = f"""<pivotCacheRecords {PIVOT_MAIN} count="4">
  <r><x v="0"/><n v="1"/></r>
  <r><x v="1"/><n v="2"/></r>
  <r><x v="0"/><n v="3"/></r>
  <r><x v="1"/><n v="4"/></r>
</pivotCacheRecords>"""

PIVOT_TABLE = f"""<pivotTableDefinition {PIVOT_MAIN} name="Totals" cacheId="1" dataCaption="Values"
    applyNumberFormats="0" applyBorderFormats="0" applyFontFormats="0"
    applyPatternFormats="0" applyAlignmentFormats="0" applyWidthHeightFormats="1"
    updatedVersion="6" minRefreshableVersion="3" createdVersion="6" indent="0"
    outline="1" outlineData="1">
  <location ref="D1:E4" firstHeaderRow="1" firstDataRow="1" firstDataCol="1"/>
  <pivotFields count="2">
    <pivotField axis="axisRow" showAll="0">
      <items count="3"><item x="0"/><item x="1"/><item t="default"/></items>
    </pivotField>
    <pivotField dataField="1" showAll="0"/>
  </pivotFields>
  <rowFields count="1"><field x="0"/></rowFields>
  <rowItems count="3"><i><x/></i><i><x v="1"/></i><i t="grand"><x/></i></rowItems>
  <colItems count="1"><i/></colItems>
  <dataFields count="1"><dataField name="Sum of Amount" fld="1" baseField="0" baseItem="0"/></dataFields>
  <pivotTableStyleInfo name="PivotStyleLight16" showRowHeaders="1" showColHeaders="1"
      showRowStripes="0" showColStripes="0" showLastColumn="1"/>
</pivotTableDefinition>"""



def pivot_table() -> Workbook:
    """A pivot table summing amounts by category, with its cache."""
    wb = _new_workbook("Pivot table")
    ws = wb.active
    ws.title = "Data"
    for row in PIVOT_DATA:
        ws.append(row)
    # The values of the pivot table as Excel shows (and stores) them.
    for coord, value in {"D1": "Row Labels", "E1": "Sum of Amount", "D2": "Fruit", "E2": 4,
                         "D3": "Veg", "E3": 6, "D4": "Grand Total", "E4": 10}.items():
        ws[coord] = value

    cache = CacheDefinition.from_tree(fromstring(PIVOT_CACHE))
    cache.records = RecordList.from_tree(fromstring(PIVOT_RECORDS))
    pivot = TableDefinition.from_tree(fromstring(PIVOT_TABLE))
    pivot.cache = cache
    ws.add_pivot(pivot)
    return wb


# ---------------------------------------------------------------------------
# Shapes: openpyxl cannot create them, so their drawing XML (as Excel writes
# it) is added to the saved file.
# ---------------------------------------------------------------------------

_DRAWING_NS = ('xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
               'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
               'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')


def _anchor(start, end, body):
    (col1, row1), (col2, row2) = start, end
    return (f'<xdr:twoCellAnchor {_DRAWING_NS}>'
            f'<xdr:from><xdr:col>{col1}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{row1}</xdr:row>'
            f'<xdr:rowOff>0</xdr:rowOff></xdr:from>'
            f'<xdr:to><xdr:col>{col2}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{row2}</xdr:row>'
            f'<xdr:rowOff>0</xdr:rowOff></xdr:to>{body}<xdr:clientData/></xdr:twoCellAnchor>')


def _shape(shape_id, name, geometry, fill, text=(), text_box=False, offset=(0, 0), size=(0, 0),
           link=None, picture=None):
    paragraphs = "".join(f'<a:p><a:r><a:rPr lang="en-US" sz="1100"/><a:t>{line}</a:t></a:r></a:p>'
                         for line in text) or '<a:p><a:endParaRPr lang="en-US" sz="1100"/></a:p>'
    text_box_attr = ' txBox="1"' if text_box else ""
    fill_xml = f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>' if fill else "<a:noFill/>"
    if picture:
        fill_xml = f'<a:blipFill><a:blip r:embed="{picture}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'
    link_xml = f'<a:hlinkClick r:id="{link}"/>' if link else ""
    return (f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{shape_id}" name="{name}">{link_xml}'
            f'</xdr:cNvPr>'
            f'<xdr:cNvSpPr{text_box_attr}/></xdr:nvSpPr>'
            f'<xdr:spPr><a:xfrm><a:off x="{offset[0]}" y="{offset[1]}"/><a:ext cx="{size[0]}" cy="{size[1]}"/>'
            f'</a:xfrm><a:prstGeom prst="{geometry}"><a:avLst/></a:prstGeom>{fill_xml}'
            f'<a:ln w="9525"><a:solidFill><a:srgbClr val="305496"/></a:solidFill></a:ln></xdr:spPr>'
            f'<xdr:txBody><a:bodyPr wrap="square" rtlCol="0" anchor="t"/><a:lstStyle/>{paragraphs}</xdr:txBody>'
            f'</xdr:sp>')


_SHAPES_XML = {
    "Diagram": [
        _anchor((0, 1), (2, 7), _shape(9, "Frame 8", "rect", "D9D9D9", text=("Behind the picture",))),
        _anchor((4, 1), (8, 6), _shape(2, "TextBox 1", "rect", None, text_box=True, text=(
            "Text box with two paragraphs.", "Shapes are kept as drawing XML."))),
        _anchor((1, 8), (3, 12), _shape(3, "Rectangle 2", "rect", "DDEBF7", text=("Start",))),
        _anchor((6, 8), (8, 12), _shape(4, "Rounded Rectangle 3", "roundRect", "E2EFDA", text=("Finish",))),
        _anchor((3, 10), (6, 10), (
            '<xdr:cxnSp macro=""><xdr:nvCxnSpPr><xdr:cNvPr id="5" name="Straight Arrow Connector 4"/>'
            '<xdr:cNvCxnSpPr><a:stCxn id="3" idx="3"/><a:endCxn id="4" idx="1"/></xdr:cNvCxnSpPr>'
            '</xdr:nvCxnSpPr><xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></a:xfrm>'
            '<a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom>'
            '<a:ln w="19050"><a:solidFill><a:srgbClr val="305496"/></a:solidFill><a:tailEnd type="triangle"/>'
            '</a:ln></xdr:spPr></xdr:cxnSp>')),
        _anchor((1, 14), (5, 19), (
            '<xdr:grpSp><xdr:nvGrpSpPr><xdr:cNvPr id="8" name="Group 7"/><xdr:cNvGrpSpPr/></xdr:nvGrpSpPr>'
            '<xdr:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2400000" cy="900000"/>'
            '<a:chOff x="0" y="0"/><a:chExt cx="2400000" cy="900000"/></a:xfrm></xdr:grpSpPr>'
            + _shape(6, "Oval 5", "ellipse", "FCE4D6", text=("A",), size=(900000, 900000))
            + _shape(7, "Oval 6", "ellipse", "FFF2CC", text=("B",), offset=(1500000, 0), size=(900000, 900000))
            + '</xdr:grpSp>')),
        _anchor((9, 1), (12, 4), _shape(10, "Link 9", "roundRect", "DDEBF7", text=("Project page",), link="rId1")),
        _anchor((9, 8), (12, 12), _shape(11, "Picture fill 10", "ellipse", None, picture="rId1")),
    ],
    "Notes": [
        _anchor((1, 1), (6, 5), _shape(2, "TextBox 1", "rect", "FFF2CC", text_box=True, text=(
            "A sheet with shapes only:", "the drawing is created on import."))),
    ],
}


# Relationships and z-order of some shapes (see xlsx2txt.shapes).
_SHAPE_EXTRAS = {
    ("Diagram", 0): {"layer": 0},
    ("Diagram", 6): {"rels": {"rId1": {"type": "hyperlink", "target": "https://github.com/click0/xlsx2txt",
                                       "external": True}}},
    ("Diagram", 7): {"rels": {"rId1": {"type": "image", "file": "fill.png"}}},
}
SHAPES = {
    sheet: [{"xml": xml, **_SHAPE_EXTRAS.get((sheet, number), {})} for number, xml in enumerate(fragments)]
    for sheet, fragments in _SHAPES_XML.items()
}


def shapes() -> tuple[Workbook, dict]:
    """Text box, shapes with text, a connector arrow, a group, a shape behind a
    picture, a shape with a hyperlink and one filled with a picture."""
    wb = _new_workbook("Shapes")
    ws = wb.active
    ws.title = "Diagram"
    ws["A1"] = "Process"
    ws.add_image(Image(io.BytesIO(_png((48, 84, 150), (48, 48)))), "A3")
    wb.create_sheet("Notes")["A1"] = "See the note"
    return wb, {"shapes": SHAPES, "media": {"fill.png": _png((237, 125, 49), (16, 16))}}


# ---------------------------------------------------------------------------
# Worksheet extensions and threaded comments, as Excel writes them.
# ---------------------------------------------------------------------------

X14 = 'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"'
XM = 'xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"'
DATA_BAR_ID = "{9F0C2E1A-5B7D-4C3E-8A21-6D4F0B9E7C15}"
SPARKLINE_COLORS = "".join(
    f'<x14:{name} rgb="{rgb}"/>' for name, rgb in [
        ("colorSeries", "FF376092"), ("colorNegative", "FFD00000"), ("colorAxis", "FF000000"),
        ("colorMarkers", "FFD00000"), ("colorFirst", "FFD00000"), ("colorLast", "FFD00000"),
        ("colorHigh", "FFD00000"), ("colorLow", "FFD00000"),
    ]
)
EXTENSIONS = [
    {"type": "conditional formatting", "xml": (
        f'<ext uri="{{78C0D931-6437-407d-A8EE-F0AAD7539E65}}" {X14}><x14:conditionalFormattings>'
        f'<x14:conditionalFormatting {XM}><x14:cfRule type="dataBar" id="{DATA_BAR_ID}">'
        '<x14:dataBar minLength="0" maxLength="100" gradient="0" negativeBarColorSameAsPositive="0" '
        'axisPosition="middle"><x14:cfvo type="autoMin"/><x14:cfvo type="autoMax"/>'
        '<x14:negativeFillColor rgb="FFFF0000"/><x14:axisColor rgb="FF000000"/></x14:dataBar></x14:cfRule>'
        '<xm:sqref>F2:F5</xm:sqref></x14:conditionalFormatting></x14:conditionalFormattings></ext>')},
    {"type": "data validation", "xml": (
        f'<ext uri="{{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}}" {X14}><x14:dataValidations count="1" {XM}>'
        '<x14:dataValidation type="list" allowBlank="1" showInputMessage="1" showErrorMessage="1">'
        '<x14:formula1><xm:f>Lists!$A$1:$A$3</xm:f></x14:formula1><xm:sqref>H2:H5</xm:sqref>'
        '</x14:dataValidation></x14:dataValidations></ext>')},
    {"type": "sparklines", "xml": (
        f'<ext uri="{{05C60535-1F16-4fd2-B633-F4F36F0B64E0}}" {X14}><x14:sparklineGroups {XM}>'
        f'<x14:sparklineGroup lineWeight="1.5" displayEmptyCellsAs="gap" markers="1">{SPARKLINE_COLORS}<x14:sparklines>'
        + "".join(f"<x14:sparkline><xm:f>Sales!B{row}:E{row}</xm:f><xm:sqref>G{row}</xm:sqref></x14:sparkline>"
                  for row in range(2, 6))
        + '</x14:sparklines></x14:sparklineGroup></x14:sparklineGroups></ext>')},
]
PERSONS = [
    {"displayName": "Olena Koval", "id": "{4B1E7C2A-0D3F-4A5B-9C6D-7E8F9A0B1C2D}",
     "userId": "olena@example.com", "providerId": "None"},
    {"displayName": "Taras Melnyk", "id": "{5C2F8D3B-1E4A-4B6C-8D7E-9F0A1B2C3D4E}",
     "userId": "taras@example.com", "providerId": "None"},
]
THREAD_ID = "{6D3A9E4C-2F5B-4C7D-9E8F-0A1B2C3D4E5F}"
THREADS = [
    {"ref": "A3", "id": THREAD_ID, "personId": PERSONS[0]["id"], "dT": "2026-01-01T12:00:00.00",
     "text": "Is Q4 final?"},
    {"ref": "A3", "id": "{7E4B0F5D-3A6C-4D8E-8F9A-1B2C3D4E5F60}", "parentId": THREAD_ID,
     "personId": PERSONS[1]["id"], "dT": "2026-01-02T09:30:00.00", "text": "Yes, checked twice."},
]
LEGACY_NOTE = (
    "[Threaded comment]\n\nYour version of Excel allows you to read this threaded comment; however, any edits "
    "to it will get removed if the file is opened in a newer version of Excel. Learn more: "
    "https://go.microsoft.com/fwlink/?linkid=870924\n\nComment:\n    Is Q4 final?\nReply:\n    Yes, checked twice."
)


# Cell metadata as Excel writes it for dynamic array formulas.
DYNAMIC_ARRAY_METADATA = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
    'xmlns:xda="http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray">'
    '<metadataTypes count="1"><metadataType name="XLDAPR" minSupportedVersion="120000" copy="1" pasteAll="1" '
    'pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" clearComments="1" assign="1" '
    'coerce="1" cellMeta="1"/></metadataTypes><futureMetadata name="XLDAPR" count="1"><bk><extLst>'
    '<ext uri="{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}"><xda:dynamicArrayProperties fDynamic="1" fCollapsed="0"/>'
    '</ext></extLst></bk></futureMetadata><cellMetadata count="1"><bk><rc t="1" v="0"/></bk></cellMetadata>'
    '</metadata>'
)


def extensions() -> tuple[Workbook, dict]:
    """Sparklines, a data bar with extended options, a drop-down list from
    another sheet, a threaded comment with a reply, a dynamic array formula
    and a switched-off error check."""
    wb = _new_workbook("Extensions")
    ws = wb.active
    ws.title = "Sales"
    _header(ws, ["Region", "Q1", "Q2", "Q3", "Q4", "Change", "Trend", "Status"])
    for row, (region, *quarters) in enumerate([
        ("North", 120, 135, 128, 160), ("South", 90, 85, 70, 65),
        ("East", 60, 75, 95, 110), ("West", 140, 120, 150, 145),
    ], start=2):
        ws.append([region, *quarters, None, None, "Open"])
        ws[f"F{row}"] = f"=E{row}-B{row}"
    for column, width in zip("ABCDEFGH", (12, 8, 8, 8, 8, 10, 16, 12)):
        ws.column_dimensions[column].width = width
    ws.conditional_formatting.add("F2:F5", DataBarRule(start_type="min", end_type="max", color="FF638EC6"))
    ws["A3"].comment = Comment(LEGACY_NOTE, f"tc={THREAD_ID}")
    ws["J1"], ws["K1"] = "Sorted", "Code"
    ws["J2"] = ArrayFormula("J2:J5", "=_xlfn._xlws.SORT(A2:A5)")
    for row, region in enumerate(sorted(["North", "South", "East", "West"]), start=2):
        if row > 2:
            ws[f"J{row}"] = region  # the spilled values Excel stores
    ws["K2"] = "0042"  # a number stored as text; its error check is switched off
    lists = wb.create_sheet("Lists")
    for value in ("Open", "In progress", "Done"):
        lists.append([value])
    return wb, {
        "sheet_extensions": {"Sales": EXTENSIONS},
        "cf_ids": {"Sales": {"F2:F5#1": DATA_BAR_ID}},
        "sheet_threads": {"Sales": THREADS},
        "persons": PERSONS,
        "cell_metadata": {"Sales": {"J2": 1}},
        "metadata": DYNAMIC_ARRAY_METADATA,
        "ignored_errors": {"Sales": '<ignoredErrors><ignoredError sqref="K2" numberStoredAsText="1"/></ignoredErrors>'},
    }


EXAMPLES = {
    "01-basic-data": basic_data,
    "02-styles": styles,
    "03-formulas": formulas,
    "04-images-charts": images_and_charts,
    "05-rich-text": rich_text,
    "06-workbook-features": workbook_features,
    "07-pivot-table": pivot_table,
    "08-shapes": shapes,
    "09-extensions": extensions,
}


# ---------------------------------------------------------------------------
# Saving
# ---------------------------------------------------------------------------

_MODIFIED = re.compile(rb"<dcterms:modified([^>]*)>[^<]*</dcterms:modified>")


def save_deterministic(wb: Workbook, path: Path, patches: dict | None = None) -> None:
    """Save reproducibly: openpyxl stamps the document and zip entries with 'now'."""
    from xlsx2txt.package import patch_package

    wb.save(path)
    if patches:
        patch_package(path, **patches)
    buffer = io.BytesIO(path.read_bytes())
    stamp = FIXED_DATE.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = _MODIFIED.sub(lambda m: b"<dcterms:modified" + m.group(1) + b">" + stamp
                                     + b"</dcterms:modified>", data)
            info = zipfile.ZipInfo(item.filename, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, data)


def generate(target_dir: Path = EXAMPLES_DIR, export: bool = True) -> list:
    """Write every example workbook (and its export) into ``target_dir``."""
    from xlsx2txt import export_xlsx

    written = []
    for name, build in EXAMPLES.items():
        path = target_dir / f"{name}.xlsx"
        built = build()
        # Builders return the workbook, or the workbook and what to add to the
        # saved file that openpyxl cannot write (see xlsx2txt.package.patch_package).
        wb, patches = built if isinstance(built, tuple) else (built, None)
        save_deterministic(wb, path, patches)
        written.append(path)
        if export:
            export_xlsx(path, target_dir / name, force=True)
    return written


if __name__ == "__main__":
    sys.path.insert(0, str(EXAMPLES_DIR.parent))
    for written_path in generate():
        print(f"{written_path.relative_to(EXAMPLES_DIR.parent)} -> {written_path.stem}/")
