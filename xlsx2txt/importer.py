"""Build an Excel workbook from the xlsx2txt model."""

import datetime
import io
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScale, DataBar, IconSet, Rule
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.packaging.relationship import Relationship
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.external_link.external import ExternalLink
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table
from openpyxl.xml.functions import fromstring

from xlsx2txt.drawings import build_image, chart_from_json
from xlsx2txt.package import patch_package
from xlsx2txt.pivots import PivotBuilder
from xlsx2txt.styles import (
    StyleApplier,
    rich_text_from_json,
    border_from_json,
    color_from_json,
    dxf_from_json,
    fill_from_json,
    font_from_json,
)

EXTERNAL_LINK_PATH = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath"

_RULE_XML_CLASSES = {"colorScale": ColorScale, "dataBar": DataBar, "iconSet": IconSet}

_ROOT_RELS = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'
)


def decode_value(value: Any, value_type: str) -> Any:
    """Inverse of :func:`xlsx2txt.exporter.encode_value`."""
    if value is None:
        return None
    if value_type == "d":
        if "T" in value:
            return datetime.datetime.fromisoformat(value)
        if ":" in value:
            return datetime.time.fromisoformat(value)
        return datetime.date.fromisoformat(value)
    if value_type == "rs":
        return rich_text_from_json(value)
    if value_type == "td":
        return datetime.timedelta(seconds=value)
    return value


def _set_attrs(obj: Any, data: dict[str, Any]) -> None:
    for name, value in data.items():
        setattr(obj, name, value)


def _apply_default_style(wb: Workbook, styles: dict[str, Any]) -> None:
    """Make the exported default style (index 0) the workbook default."""
    if not styles.get("cellStyles"):
        return
    default = styles["cellStyles"][0]
    wb._fonts = IndexedList([font_from_json(styles["fonts"][default["font"]])])
    fills = list(wb._fills)
    fills[0] = fill_from_json(styles["fills"][default["fill"]])
    wb._fills = IndexedList(fills)
    wb._borders = IndexedList([border_from_json(styles["borders"][default["border"]])])


def _write_cell(ws, coord: str, record: dict[str, Any], applier: StyleApplier) -> None:
    column_letter, row = coordinate_from_string(coord)
    cell = ws.cell(row=row, column=column_index_from_string(column_letter))

    if "s" in record:
        applier.apply(cell, record["s"])

    if isinstance(cell, MergedCell):
        return

    if "link" in record:
        cell.hyperlink = Hyperlink(ref=coord, **record["link"])

    value_type = record.get("t")
    if value_type == "f":
        f_type = record.get("fType")
        if f_type == "array":
            cell.value = ArrayFormula(ref=record["fRef"], text=record["f"])
        elif f_type == "dataTable":
            cell.value = DataTableFormula(**record["fAttrs"])
        else:
            cell.value = record["f"]
    else:
        cell.value = decode_value(record.get("v"), value_type)
        # Keep literal strings that look like formulas or errors as strings.
        if value_type == "s" and cell.data_type != "s":
            cell.data_type = "s"

    if "comment" in record:
        comment = record["comment"]
        cell.comment = Comment(comment.get("text", ""), comment.get("author") or "")


def _import_dimensions(ws, dims: dict[str, Any], applier: StyleApplier) -> None:
    for letter, data in dims.get("columns", {}).items():
        dim = ws.column_dimensions[letter]
        if "width" in data:
            dim.width = data["width"]
        dim.hidden = data.get("hidden", False)
        dim.outline_level = data.get("outlineLevel", 0)
        dim.collapsed = data.get("collapsed", False)
        dim.bestFit = data.get("bestFit", False)
        if "min" in data:
            dim.min = data["min"]
            dim.max = data["max"]
        if "s" in data:
            applier.apply(dim, data["s"])

    for number, data in dims.get("rows", {}).items():
        dim = ws.row_dimensions[int(number)]
        if "height" in data:
            dim.height = data["height"]
        dim.hidden = data.get("hidden", False)
        dim.outline_level = data.get("outlineLevel", 0)
        dim.collapsed = data.get("collapsed", False)
        if "s" in data:
            applier.apply(dim, data["s"])


def _import_sheet(wb: Workbook, sheet: dict[str, Any], applier: StyleApplier,
                  media: dict[str, bytes], pivots: PivotBuilder) -> None:
    ws = wb.create_sheet(sheet["name"])
    ws.sheet_state = sheet.get("state", "visible")

    if "tabColor" in sheet:
        ws.sheet_properties.tabColor = color_from_json(sheet["tabColor"])

    view = dict(sheet.get("view", {}))
    freeze = view.pop("freezePanes", None)
    _set_attrs(ws.sheet_view, view)
    if freeze:
        ws.freeze_panes = freeze

    _set_attrs(ws.sheet_format, sheet.get("format", {}))
    _import_dimensions(ws, sheet.get("dimensions", {}), applier)

    for rng in sheet.get("mergedCells", []):
        ws.merge_cells(rng)

    for coord, record in sheet.get("cells", {}).items():
        _write_cell(ws, coord, record, applier)

    if sheet.get("autoFilter"):
        ws.auto_filter.ref = sheet["autoFilter"]

    printing = sheet.get("print", {})
    if "area" in printing:
        ws.print_area = printing["area"]
    if "titleRows" in printing:
        ws.print_title_rows = printing["titleRows"]
    if "titleCols" in printing:
        ws.print_title_cols = printing["titleCols"]
    _set_attrs(ws.page_setup, printing.get("pageSetup", {}))
    _set_attrs(ws.page_margins, printing.get("margins", {}))
    _set_attrs(ws.print_options, printing.get("options", {}))
    if printing.get("fitToPage"):
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    if "protection" in sheet:
        protection = dict(sheet["protection"])
        password = protection.pop("password", None)
        _set_attrs(ws.protection, protection)
        if password:
            ws.protection.set_password(password, already_hashed=True)

    for data in sheet.get("dataValidations", []):
        data = dict(data)
        sqref = data.pop("sqref")
        dv = DataValidation(**data)
        dv.sqref = sqref
        ws.add_data_validation(dv)

    for block in sheet.get("conditionalFormatting", []):
        for data in block["rules"]:
            data = dict(data)
            kwargs = {}
            for part, cls in _RULE_XML_CLASSES.items():
                if part in data:
                    kwargs[part] = cls.from_tree(fromstring(data.pop(part)))
            dxf = dxf_from_json(data.pop("dxf", None))
            rule = Rule(dxf=dxf, **data, **kwargs)
            ws.conditional_formatting.add(block["range"], rule)

    for image in sheet.get("images", []):
        if image["file"] not in media:
            raise ValueError(f"Sheet '{sheet['name']}': missing image file media/{image['file']}")
        ws.add_image(build_image(media[image["file"]], image["anchor"]))

    for chart in sheet.get("charts", []):
        ws.add_chart(chart_from_json(chart))

    for entry in sheet.get("pivotTables", []):
        ws.add_pivot(pivots.table(entry))

    for xml in sheet.get("tables", []):
        ws.add_table(Table.from_tree(fromstring(xml)))

    for data in sheet.get("definedNames", []):
        ws.defined_names[data["name"]] = _defined_name(data)


def _defined_name(data: dict[str, Any]) -> DefinedName:
    return DefinedName(
        name=data["name"],
        attr_text=data["value"],
        comment=data.get("comment"),
        hidden=data.get("hidden"),
    )


_CONTENT_TYPES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Override PartName="/xl/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/>'
    '</Types>'
)


def _custom_property(data: dict[str, Any]):
    from openpyxl.packaging import custom

    cls = getattr(custom, data.get("type", ""), None)
    if cls is None or not data.get("type", "").endswith("Property"):
        raise ValueError(f"Unknown custom property type: {data.get('type')!r}")
    value = data["value"]
    if cls is custom.DateTimeProperty:
        value = datetime.datetime.fromisoformat(value)
    return cls(name=data["name"], value=value)


def _external_link(data: dict[str, Any]) -> ExternalLink:
    link = ExternalLink.from_tree(fromstring(data["xml"]))
    link.file_link = Relationship(
        Id="rId1",
        Type=data.get("type", EXTERNAL_LINK_PATH),
        Target=data["target"],
        TargetMode=data.get("targetMode", "External"),
    )
    return link


def _vba_archive(parts: dict[str, bytes]) -> zipfile.ZipFile:
    """Build an in-memory archive that openpyxl merges into the saved file."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", _CONTENT_TYPES)
        if "_rels/.rels" not in parts:
            archive.writestr("_rels/.rels", _ROOT_RELS)
        for name, content in parts.items():
            archive.writestr(name, content)
    buffer.seek(0)
    return zipfile.ZipFile(buffer, "r")


def build_workbook(model: dict[str, Any]) -> Workbook:
    """Create an openpyxl workbook from a model."""
    wb = Workbook()
    wb.remove(wb.active)

    workbook = model["workbook"]
    wb.epoch = CALENDAR_MAC_1904 if workbook.get("epoch") == 1904 else CALENDAR_WINDOWS_1900

    styles = model["styles"]
    _apply_default_style(wb, styles)
    applier = StyleApplier(styles)
    pivots = PivotBuilder(model.get("pivots") or {})

    for sheet in model["sheets"]:
        _import_sheet(wb, sheet, applier, model.get("media") or {}, pivots)

    for data in sorted(workbook.get("chartsheets", []), key=lambda d: d["position"]):
        chartsheet = wb.create_chartsheet(data["name"], data["position"])
        chartsheet.sheet_state = data.get("state", "visible")
        for chart in data.get("charts", []):
            chartsheet.add_chart(chart_from_json(chart))

    for data in workbook.get("customProperties", []):
        wb.custom_doc_props.append(_custom_property(data))

    for data in workbook.get("externalLinks", []):
        wb._external_links.append(_external_link(data))

    if wb.worksheets:
        active = workbook.get("activeSheet", 0)
        active = active if 0 <= active < len(wb.worksheets) else 0
        wb.active = wb._sheets.index(wb.worksheets[active])

    for name, value in workbook.get("properties", {}).items():
        if name in ("created", "modified", "lastPrinted"):
            value = datetime.datetime.fromisoformat(value)
        setattr(wb.properties, name, value)

    _set_attrs(wb.calculation, workbook.get("calculation", {}))

    for data in workbook.get("definedNames", []):
        wb.defined_names[data["name"]] = _defined_name(data)

    if model.get("theme"):
        wb.loaded_theme = model["theme"].encode("utf-8")

    if model.get("vba"):
        wb.vba_archive = _vba_archive(model["vba"])

    return wb


def import_model(model: dict[str, Any], output: str | Path) -> Path:
    """Build a workbook from a model and save it to ``output``."""
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(model).save(output)
    # openpyxl does not write printer settings and shapes; add them to the saved package.
    files = model.get("printerSettings") or {}
    settings = {
        sheet["name"]: files[sheet["printerSettings"]]
        for sheet in model["sheets"]
        if sheet.get("printerSettings") in files
    }
    shapes = {sheet["name"]: sheet.get("shapes", []) for sheet in model["sheets"]}
    patch_package(output, printer_settings=settings, shapes=shapes, media=model.get("media") or {})
    return output
