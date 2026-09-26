"""Direct work with the xlsx package (zip) for parts openpyxl does not keep.

- Printer settings (``xl/printerSettings/*.bin``) are read from the source
  archive and written back into the saved file.
- Shapes are read from the drawings and written back into them
  (see :mod:`xlsx2txt.shapes`).
- Parts and sheet features that openpyxl drops are detected so that export
  can warn about them instead of losing them silently.
"""

import hashlib
import html
import posixpath
import re
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl.packaging.relationship import get_dependents, get_rels_path
from openpyxl.reader.workbook import WorkbookParser

from xlsx2txt.shapes import (
    EMPTY_DRAWING,
    add_to_drawing,
    classify,
    drawing_children,
    full_type,
    rewrite_rel_ids,
    shapes_from_drawing,
)

REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PRINTER_SETTINGS_REL = f"{REL_NS}/printerSettings"
PRINTER_SETTINGS_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.printerSettings"
DRAWING_REL = f"{REL_NS}/drawing"
DRAWING_TYPE = "application/vnd.openxmlformats-officedocument.drawing+xml"


def _sheet_paths(archive: zipfile.ZipFile) -> dict[str, str]:
    """Map sheet (and chart sheet) names to their part names."""
    parser = WorkbookParser(archive, "xl/workbook.xml")
    parser.parse()
    return {sheet.name: rel.target for sheet, rel in parser.find_sheets()}


def _sheet_drawings(archive: zipfile.ZipFile, names: set[str]) -> dict[str, list[str]]:
    """Map sheet (and chart sheet) names to the drawing parts they use."""
    result: dict[str, list[str]] = {}
    for sheet_name, sheet_path in _sheet_paths(archive).items():
        rels_path = get_rels_path(sheet_path)
        if sheet_path not in names or rels_path not in names:
            continue
        drawings = [rel.target for rel in get_dependents(archive, rels_path).find(DRAWING_REL)
                    if rel.target in names]
        if drawings:
            result[sheet_name] = drawings
    return result


# ---------------------------------------------------------------------------
# Printer settings
# ---------------------------------------------------------------------------

def printer_settings_name(content: bytes) -> str:
    """Content-addressed file name; sheets often share identical settings."""
    return f"printer_{hashlib.sha256(content).hexdigest()[:16]}.bin"


def extract_printer_settings(path) -> dict[str, bytes]:
    """Return ``{sheet name: printer settings bytes}`` for worksheets that have them."""
    result: dict[str, bytes] = {}
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        for sheet_name, sheet_path in _sheet_paths(archive).items():
            rels_path = get_rels_path(sheet_path)
            if rels_path not in names or not sheet_path.startswith("xl/worksheets/"):
                continue
            for rel in get_dependents(archive, rels_path).find(PRINTER_SETTINGS_REL):
                if rel.target in names:
                    result[sheet_name] = archive.read(rel.target)
                    break
    return result


_RELATIONSHIPS_EMPTY = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>'
)


def _add_relationship(rels_xml: str, rel_type: str, target: str, external: bool = False) -> tuple[str, str]:
    used = set(re.findall(r'\bId="([^"]+)"', rels_xml))
    number = 1
    while f"rId{number}" in used:
        number += 1
    rel_id = f"rId{number}"
    mode = ' TargetMode="External"' if external else ""
    element = f'<Relationship Id="{rel_id}" Type="{rel_type}" Target="{html.escape(target)}"{mode}/>'
    return rels_xml.replace("</Relationships>", element + "</Relationships>", 1), rel_id


def _declare_r(sheet_xml: str) -> str:
    root = re.search(r"<worksheet\b[^>]*>", sheet_xml)
    if root and "xmlns:r=" not in root.group(0):
        new_root = root.group(0)[:-1] + f' xmlns:r="{REL_NS}">'
        sheet_xml = sheet_xml.replace(root.group(0), new_root, 1)
    return sheet_xml


def _link_page_setup(sheet_xml: str, rel_id: str) -> str:
    """Point <pageSetup r:id=...> at the printer settings relationship."""
    sheet_xml = _declare_r(sheet_xml)
    setup = re.search(r"<pageSetup\b[^>]*?(/?)>", sheet_xml)
    if setup:
        tag = setup.group(0)
        tag = re.sub(r'\s+r:id="[^"]*"', "", tag)
        tag = re.sub(r"\s*(/?)>$", rf' r:id="{rel_id}"\1>', tag)
        return sheet_xml[:setup.start()] + tag + sheet_xml[setup.end():]
    margins = re.search(r"<pageMargins\b[^>]*/>", sheet_xml)
    if not margins:
        raise ValueError("worksheet without <pageMargins>: cannot attach printer settings")
    element = f'<pageSetup r:id="{rel_id}"/>'
    return sheet_xml[:margins.end()] + element + sheet_xml[margins.end():]


# Elements that follow <drawing> in a worksheet (ECMA-376, CT_Worksheet).
_AFTER_DRAWING = re.compile(
    r"<(legacyDrawing|legacyDrawingHF|drawingHF|picture|oleObjects|controls|webPublishItems|tableParts|extLst)\b"
    r"|</worksheet>"
)


def _link_drawing(sheet_xml: str, rel_id: str) -> str:
    sheet_xml = _declare_r(sheet_xml)
    position = _AFTER_DRAWING.search(sheet_xml).start()
    return sheet_xml[:position] + f'<drawing r:id="{rel_id}"/>' + sheet_xml[position:]


def _free_part(parts: dict[str, bytes], pattern: str) -> str:
    number = 1
    while pattern.format(number) in parts:
        number += 1
    return pattern.format(number)


def _add_override(parts: dict[str, bytes], part: str, content_type: str) -> None:
    types = parts["[Content_Types].xml"].decode("utf-8")
    override = f'<Override PartName="/{part}" ContentType="{content_type}"/>'
    parts["[Content_Types].xml"] = types.replace("</Types>", override + "</Types>", 1).encode("utf-8")


def _add_sheet_relationship(parts: dict[str, bytes], sheet_path: str, rel_type: str, part: str) -> str:
    rels_path = get_rels_path(sheet_path)
    rels_xml = parts[rels_path].decode("utf-8") if rels_path in parts else _RELATIONSHIPS_EMPTY
    target = posixpath.relpath(part, posixpath.dirname(sheet_path))
    rels_xml, rel_id = _add_relationship(rels_xml, rel_type, target)
    parts[rels_path] = rels_xml.encode("utf-8")
    return rel_id


def _add_printer_settings(parts, sheet_paths, settings: dict[str, bytes]) -> None:
    for sheet_name, content in sorted(settings.items()):
        sheet_path = sheet_paths.get(sheet_name)
        if sheet_path is None:
            raise ValueError(f"no sheet named {sheet_name!r} for printer settings")
        part = _free_part(parts, "xl/printerSettings/printerSettings{}.bin")
        parts[part] = content
        rel_id = _add_sheet_relationship(parts, sheet_path, PRINTER_SETTINGS_REL, part)
        parts[sheet_path] = _link_page_setup(parts[sheet_path].decode("utf-8"), rel_id).encode("utf-8")
        _add_override(parts, part, PRINTER_SETTINGS_TYPE)


_IMAGE_TYPES = {
    "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "gif": "image/gif", "bmp": "image/bmp",
    "tif": "image/tiff", "tiff": "image/tiff", "emf": "image/x-emf", "wmf": "image/x-wmf",
    "svg": "image/svg+xml", "wdp": "image/vnd.ms-photo",
}


def _ensure_default(parts: dict[str, bytes], extension: str) -> None:
    types = parts["[Content_Types].xml"].decode("utf-8")
    if re.search(rf'<Default\b[^>]*Extension="{re.escape(extension)}"', types, re.I):
        return
    content_type = _IMAGE_TYPES.get(extension.lower(), "application/octet-stream")
    default = f'<Default Extension="{extension}" ContentType="{content_type}"/>'
    types = re.sub(r"(<Types\b[^>]*>)", lambda m: m.group(1) + default, types, count=1)
    parts["[Content_Types].xml"] = types.encode("utf-8")


def _add_shapes(parts, sheet_paths, drawings, shapes: dict[str, list[dict]],
                media: dict[str, bytes]) -> None:
    media_parts: dict[str, str] = {}
    for sheet_name, sheet_shapes in sorted(shapes.items()):
        if not sheet_shapes:
            continue
        sheet_path = sheet_paths.get(sheet_name)
        if sheet_path is None:
            raise ValueError(f"no sheet named {sheet_name!r} for shapes")
        if sheet_name in drawings:
            part = drawings[sheet_name][0]
        else:
            part = _free_part(parts, "xl/drawings/drawing{}.xml")
            parts[part] = EMPTY_DRAWING.encode("utf-8")
            rel_id = _add_sheet_relationship(parts, sheet_path, DRAWING_REL, part)
            parts[sheet_path] = _link_drawing(parts[sheet_path].decode("utf-8"), rel_id).encode("utf-8")
            _add_override(parts, part, DRAWING_TYPE)

        rels_path = get_rels_path(part)
        rels_xml = parts[rels_path].decode("utf-8") if rels_path in parts else _RELATIONSHIPS_EMPTY
        placed = []
        for shape in sheet_shapes:
            mapping = {}
            for old_id, rel in (shape.get("rels") or {}).items():
                if rel.get("external"):
                    target, external = rel["target"], True
                else:
                    name = rel["file"]
                    if name not in media:
                        raise ValueError(f"Sheet '{sheet_name}': missing image file media/{name} of a shape")
                    if name not in media_parts:
                        extension = posixpath.splitext(name)[1].lstrip(".") or "png"
                        media_part = _free_part(parts, "xl/media/image{}." + extension)
                        parts[media_part] = media[name]
                        _ensure_default(parts, extension)
                        media_parts[name] = media_part
                    target = posixpath.relpath(media_parts[name], posixpath.dirname(part))
                    external = False
                rels_xml, mapping[old_id] = _add_relationship(rels_xml, full_type(rel["type"]), target, external)
            placed.append((rewrite_rel_ids(shape["xml"], mapping), shape.get("layer")))
        if rels_xml != _RELATIONSHIPS_EMPTY:
            parts[rels_path] = rels_xml.encode("utf-8")
        parts[part] = add_to_drawing(parts[part].decode("utf-8"), placed).encode("utf-8")


def patch_package(path, printer_settings: dict[str, bytes] | None = None,
                  shapes: dict[str, list[dict]] | None = None, media: dict[str, bytes] | None = None) -> None:
    """Add printer settings and shapes to a file saved by openpyxl (in place).

    ``printer_settings``: ``{sheet name: bytes}``; ``shapes``: ``{sheet name:
    [shape, ...]}`` as exported (see :mod:`xlsx2txt.shapes`); ``media``: the
    pictures shapes refer to.
    """
    if not printer_settings and not any((shapes or {}).values()):
        return
    path = Path(path)
    with zipfile.ZipFile(path) as source:
        parts = {info.filename: source.read(info.filename) for info in source.infolist()}
        sheet_paths = _sheet_paths(source)
        drawings = _sheet_drawings(source, set(parts))

    _add_printer_settings(parts, sheet_paths, printer_settings or {})
    _add_shapes(parts, sheet_paths, drawings, shapes or {}, media or {})

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target_zip:
        # [Content_Types].xml first, as Office writes it.
        for name in ["[Content_Types].xml"] + [n for n in parts if n != "[Content_Types].xml"]:
            target_zip.writestr(name, parts[name])


def inject_printer_settings(path, settings: dict[str, bytes]) -> None:
    """Add printer settings to a file saved by openpyxl (in place)."""
    patch_package(path, printer_settings=settings)


def _drawing_shapes(archive: zipfile.ZipFile, names: set[str], drawing: str):
    rels_path = get_rels_path(drawing)
    rels = {rel.Id: rel for rel in get_dependents(archive, rels_path)} if rels_path in names else {}
    return shapes_from_drawing(archive.read(drawing).decode("utf-8", errors="replace"), rels, archive.read)


def extract_shapes(path) -> tuple[dict[str, list[dict]], dict[str, bytes]]:
    """Return ``({worksheet name: [shape, ...]}, {media name: bytes})``."""
    result: dict[str, list[dict]] = {}
    media: dict[str, bytes] = {}
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        sheet_paths = _sheet_paths(archive)
        for sheet_name, drawings in _sheet_drawings(archive, names).items():
            if not sheet_paths[sheet_name].startswith("xl/worksheets/"):
                continue
            for drawing in drawings:
                shapes, _, shape_media = _drawing_shapes(archive, names, drawing)
                result.setdefault(sheet_name, []).extend(shapes)
                media.update(shape_media)
    return result, media


# ---------------------------------------------------------------------------
# Parts and features that are not exported
# ---------------------------------------------------------------------------

# Parts xlsx2txt exports or that openpyxl regenerates from exported data.
_KNOWN_PARTS = re.compile(
    r"^(\[Content_Types\]\.xml|_rels/\.rels|docProps/(core|app|custom)\.xml"
    r"|xl/(workbook\.xml|_rels/workbook\.xml\.rels|styles\.xml|sharedStrings\.xml|calcChain\.xml"
    r"|theme/.*|worksheets/.*|chartsheets/.*|drawings/.*|comments\d*\.xml|comments/.*|media/.*|charts/.*"
    r"|tables/.*|pivotTables/.*|pivotCache/.*|externalLinks/.*|printerSettings/.*"
    r"|vbaProject\.bin|vbaProjectSignature.*\.bin|_rels/vbaProject\.bin\.rels))$"
)

# Recognised groups of parts that are lost, with a readable description.
_LOST_PARTS = [
    (re.compile(r"^xl/threadedComments/"), "threaded comments (kept only as plain notes)"),
    (re.compile(r"^xl/persons/"), None),  # authors of threaded comments, reported with them
    (re.compile(r"^xl/(slicers|slicerCaches)/"), "slicers"),
    (re.compile(r"^xl/timelines?|^xl/timelineCaches/"), "timelines"),
    (re.compile(r"^xl/(ctrlProps|activeX)/"), "form controls / ActiveX controls"),
    (re.compile(r"^xl/embeddings/"), "embedded objects (OLE)"),
    (re.compile(r"^xl/(connections\.xml|queryTables/)"), "data connections / queries"),
    (re.compile(r"^customXml/"), "custom XML parts (e.g. Power Query)"),
    (re.compile(r"^xl/model/"), "data model (Power Pivot)"),
    (re.compile(r"^xl/richData/"), "pictures in cells / rich data types"),
    (re.compile(r"^xl/metadata\.xml$"), "cell metadata (dynamic array formulas may become legacy arrays)"),
    (re.compile(r"^xl/webextensions/"), "Office add-ins"),
    (re.compile(r"^customUI"), "ribbon customisation"),
]

# Sheet XML fragments of features openpyxl drops.
_SHEET_FEATURES = [
    (re.compile(r"<(\w+:)?sparklineGroups?\b"), "sparklines"),
    (re.compile(r"<x14:conditionalFormattings\b"), "extended conditional formatting (e.g. icon/data bar options)"),
    (re.compile(r"<x14:dataValidations\b"), "extended data validation (lists from other sheets)"),
    (re.compile(r"<(\w+:)?slicerList\b"), "slicers"),
]
_SMART_ART = re.compile(r"drawingml/2006/diagram")


def unsupported_warnings(path, keep_vba: bool = False) -> list[str]:
    """Describe everything in the file that the export does not keep.

    ``keep_vba``: the VBA project and ribbon customisation are exported
    (macro-enabled workbooks).
    """
    warnings: list[str] = []
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        lost = Counter()
        unknown = []
        for name in names:
            # Folders and relationship files are covered by the parts they link.
            if name.endswith("/") or name.endswith(".rels") or _KNOWN_PARTS.match(name):
                continue
            match = next(((pattern, text) for pattern, text in _LOST_PARTS if pattern.match(name)), None)
            if match is None:
                unknown.append(name)
            elif match[1] and not (keep_vba and match[1] == "ribbon customisation"):
                lost[match[1]] += 1
        for description, count in sorted(lost.items()):
            warnings.append(f"{description}: {count} part(s) are not exported")
        if unknown:
            shown = ", ".join(unknown[:5]) + (" …" if len(unknown) > 5 else "")
            warnings.append(f"{len(unknown)} unknown part(s) are not exported: {shown}")

        drawings = _sheet_drawings(archive, set(names))
        for sheet_name, sheet_path in _sheet_paths(archive).items():
            if sheet_path not in names:
                continue
            xml = archive.read(sheet_path).decode("utf-8", errors="replace")
            for pattern, description in _SHEET_FEATURES:
                if pattern.search(xml):
                    warnings.append(f"Sheet '{sheet_name}': {description} are not exported")
            shapes = 0
            smart_art = False
            is_worksheet = sheet_path.startswith("xl/worksheets/")
            for drawing_path in drawings.get(sheet_name, []):
                drawing = archive.read(drawing_path).decode("utf-8", errors="replace")
                if is_worksheet:
                    shapes += _drawing_shapes(archive, set(names), drawing_path)[1]
                else:
                    shapes += sum(1 for child in drawing_children(drawing)[1] if classify(child))
                smart_art = smart_art or bool(_SMART_ART.search(drawing))
            if shapes:
                what = "form control(s) or shape(s) with unsupported links" if is_worksheet else "shape(s)"
                warnings.append(f"Sheet '{sheet_name}': {shapes} {what} are not exported")
            if smart_art:
                warnings.append(f"Sheet '{sheet_name}': SmartArt diagrams are not exported")
    return warnings
