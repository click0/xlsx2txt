"""Export an Excel workbook into the xlsx2txt model (a JSON-friendly dict)."""

import datetime
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell as OpenpyxlCell, MergedCell
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.xml.functions import tostring

from xlsx2txt import __version__
from xlsx2txt.pivots import export_pivots
from xlsx2txt.drawings import anchor_to_json, chart_to_json, extract_images, media_name
from xlsx2txt.styles import StyleTable, color_to_json, dxf_to_json, rich_text_to_json
from xlsx2txt.vba import extract_vba_sources

FORMAT_NAME = "xlsx2txt"
FORMAT_VERSION = 1

# openpyxl attribute names copied verbatim for various sheet-level objects.
SHEET_FORMAT_ATTRS = [
    "baseColWidth", "defaultColWidth", "defaultRowHeight", "customHeight",
    "zeroHeight", "thickTop", "thickBottom", "outlineLevelRow", "outlineLevelCol",
]
SHEET_VIEW_ATTRS = [
    "zoomScale", "showGridLines", "showRowColHeaders", "showZeros",
    "rightToLeft", "tabSelected", "view",
]
PAGE_SETUP_ATTRS = [
    "orientation", "paperSize", "scale", "fitToWidth", "fitToHeight",
    "firstPageNumber", "useFirstPageNumber", "pageOrder", "blackAndWhite",
    "draft", "horizontalDpi", "verticalDpi",
]
PAGE_MARGIN_ATTRS = ["left", "right", "top", "bottom", "header", "footer"]
PRINT_OPTION_ATTRS = ["horizontalCentered", "verticalCentered", "headings", "gridLines"]
PROTECTION_ATTRS = [
    "sheet", "objects", "scenarios", "formatCells", "formatRows", "formatColumns",
    "insertColumns", "insertRows", "insertHyperlinks", "deleteColumns", "deleteRows",
    "selectLockedCells", "selectUnlockedCells", "sort", "autoFilter", "pivotTables",
    "password", "algorithmName", "hashValue", "saltValue", "spinCount",
]
DATA_VALIDATION_ATTRS = [
    "type", "operator", "formula1", "formula2", "allow_blank", "showDropDown",
    "showErrorMessage", "showInputMessage", "errorStyle", "errorTitle", "error",
    "promptTitle", "prompt", "imeMode",
]
RULE_ATTRS = [
    "type", "priority", "stopIfTrue", "aboveAverage", "percent", "bottom",
    "operator", "text", "timePeriod", "rank", "stdDev", "equalAverage",
]
RULE_XML_PARTS = ["colorScale", "dataBar", "iconSet"]
PROPERTY_ATTRS = [
    "title", "subject", "creator", "keywords", "description", "category",
    "lastModifiedBy", "language", "identifier", "version", "revision",
    "contentStatus", "created", "modified", "lastPrinted",
]
CALC_ATTRS = ["calcMode", "fullCalcOnLoad", "iterate", "iterateCount", "iterateDelta", "refMode"]


def pick_attrs(obj: Any, names: List[str], skip_defaults: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """Collect JSON-friendly attribute values of an object, skipping None."""
    result = {}
    if obj is None:
        return result
    for name in names:
        value = getattr(obj, name, None)
        if value is None:
            continue
        if skip_defaults and name in skip_defaults and skip_defaults[name] == value:
            continue
        if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
            value = value.isoformat()
        if not isinstance(value, (str, int, float, bool)):
            continue
        result[name] = value
    return result


# ---------------------------------------------------------------------------
# Cell values
# ---------------------------------------------------------------------------

def encode_value(value: Any) -> Dict[str, Any]:
    """Encode a plain (non-formula) cell value as {"v": ..., "t": ...}."""
    if isinstance(value, CellRichText):
        return {"v": rich_text_to_json(value), "t": "rs"}
    if isinstance(value, bool):
        return {"v": value, "t": "b"}
    if isinstance(value, (int, float)):
        return {"v": value, "t": "n"}
    if isinstance(value, datetime.datetime):
        return {"v": value.isoformat(), "t": "d"}
    if isinstance(value, datetime.date):
        return {"v": value.isoformat(), "t": "d"}
    if isinstance(value, datetime.time):
        return {"v": value.isoformat(), "t": "d"}
    if isinstance(value, datetime.timedelta):
        return {"v": value.total_seconds(), "t": "td"}
    return {"v": str(value), "t": "s"}


def _encode_cell(cell: OpenpyxlCell, cached: Any, styles: StyleTable) -> Optional[Dict[str, Any]]:
    record: Dict[str, Any] = {}
    value = cell.value

    if isinstance(cell, MergedCell):
        pass
    elif cell.data_type == "f":
        record["t"] = "f"
        if isinstance(value, ArrayFormula):
            record["f"] = value.text
            record["fType"] = "array"
            record["fRef"] = value.ref
        elif isinstance(value, DataTableFormula):
            record["fType"] = "dataTable"
            record["fAttrs"] = {
                k: getattr(value, k)
                for k in ("ref", "ca", "dt2D", "dtr", "r1", "r2", "del1", "del2")
                if getattr(value, k)
            }
        else:
            record["f"] = value
        if cached is not None:
            encoded = encode_value(cached)
            record["v"] = encoded["v"]
            if encoded["t"] != "n":
                record["vt"] = encoded["t"]
    elif cell.data_type == "e":
        record = {"v": value, "t": "e"}
    elif value is not None:
        record = encode_value(value)

    style_index = styles.add(cell) if cell.has_style else 0
    if style_index:
        record["s"] = style_index

    if cell.hyperlink is not None:
        link = pick_attrs(cell.hyperlink, ["target", "location", "tooltip", "display"])
        if link:
            record["link"] = link

    if cell.comment is not None:
        record["comment"] = {"text": cell.comment.text, "author": cell.comment.author}

    return record or None


# ---------------------------------------------------------------------------
# Sheets
# ---------------------------------------------------------------------------

def _export_dimensions(ws, styles: StyleTable) -> Dict[str, Any]:
    columns = {}
    for letter, dim in sorted(ws.column_dimensions.items(), key=lambda kv: kv[1].min or 0):
        data: Dict[str, Any] = {}
        if dim.width is not None:  # 0 is a real width (hidden columns)
            data["width"] = dim.width
        if dim.hidden:
            data["hidden"] = True
        if dim.outline_level:
            data["outlineLevel"] = dim.outline_level
        if dim.collapsed:
            data["collapsed"] = True
        if dim.bestFit:
            data["bestFit"] = True
        if dim.has_style:
            index = styles.add(dim)
            if index:
                data["s"] = index
        if not data:
            continue
        if dim.min and dim.max and dim.max != dim.min:
            data["min"] = dim.min
            data["max"] = dim.max
        columns[letter] = data

    rows = {}
    for number, dim in sorted(ws.row_dimensions.items()):
        data = {}
        if dim.height is not None:
            data["height"] = dim.height
        if dim.hidden:
            data["hidden"] = True
        if dim.outline_level:
            data["outlineLevel"] = dim.outline_level
        if dim.collapsed:
            data["collapsed"] = True
        if dim.has_style:
            index = styles.add(dim)
            if index:
                data["s"] = index
        if data:
            rows[str(number)] = data

    return {"columns": columns, "rows": rows}


def _export_conditional_formatting(ws) -> List[Dict[str, Any]]:
    result = []
    for cf in ws.conditional_formatting:
        rules = []
        for rule in cf.rules:
            data = pick_attrs(rule, RULE_ATTRS)
            if rule.formula:
                data["formula"] = list(rule.formula)
            for part in RULE_XML_PARTS:
                obj = getattr(rule, part)
                if obj is not None:
                    data[part] = tostring(obj.to_tree()).decode("utf-8")
            dxf = dxf_to_json(rule.dxf)
            if dxf:
                data["dxf"] = dxf
            rules.append(data)
        result.append({"range": str(cf.sqref), "rules": rules})
    return result


def _export_data_validations(ws) -> List[Dict[str, Any]]:
    result = []
    for dv in ws.data_validations.dataValidation:
        data = {"sqref": str(dv.sqref)}
        data.update(pick_attrs(dv, DATA_VALIDATION_ATTRS))
        result.append(data)
    return result


def _export_sheet(ws, ws_values, styles: StyleTable) -> Dict[str, Any]:
    sheet: Dict[str, Any] = {"name": ws.title}
    if ws.sheet_state != "visible":
        sheet["state"] = ws.sheet_state

    tab_color = color_to_json(ws.sheet_properties.tabColor)
    if tab_color is not None:
        sheet["tabColor"] = tab_color

    view = pick_attrs(ws.sheet_view, SHEET_VIEW_ATTRS)
    if ws.freeze_panes:
        view["freezePanes"] = ws.freeze_panes
    if view:
        sheet["view"] = view

    sheet_format = pick_attrs(ws.sheet_format, SHEET_FORMAT_ATTRS)
    if sheet_format:
        sheet["format"] = sheet_format

    sheet["usedRange"] = ws.dimensions
    sheet["dimensions"] = _export_dimensions(ws, styles)
    sheet["mergedCells"] = [str(rng) for rng in ws.merged_cells.ranges]

    if ws.auto_filter.ref:
        sheet["autoFilter"] = ws.auto_filter.ref

    printing: Dict[str, Any] = {}
    if ws.print_area:
        printing["area"] = ws.print_area
    if ws.print_title_rows:
        printing["titleRows"] = ws.print_title_rows
    if ws.print_title_cols:
        printing["titleCols"] = ws.print_title_cols
    page_setup = pick_attrs(ws.page_setup, PAGE_SETUP_ATTRS)
    if page_setup:
        printing["pageSetup"] = page_setup
    margins = pick_attrs(ws.page_margins, PAGE_MARGIN_ATTRS)
    if margins:
        printing["margins"] = margins
    options = pick_attrs(ws.print_options, PRINT_OPTION_ATTRS)
    options = {k: v for k, v in options.items() if v}
    if options:
        printing["options"] = options
    if ws.sheet_properties.pageSetUpPr is not None and ws.sheet_properties.pageSetUpPr.fitToPage:
        printing["fitToPage"] = True
    if printing:
        sheet["print"] = printing

    if ws.protection.sheet:
        sheet["protection"] = pick_attrs(ws.protection, PROTECTION_ATTRS)

    validations = _export_data_validations(ws)
    if validations:
        sheet["dataValidations"] = validations

    conditional = _export_conditional_formatting(ws)
    if conditional:
        sheet["conditionalFormatting"] = conditional

    if ws.tables:
        sheet["tables"] = [
            tostring(table.to_tree()).decode("utf-8") for table in ws.tables.values()
        ]

    local_names = _export_defined_names(ws.defined_names.values())
    if local_names:
        sheet["definedNames"] = local_names

    cells = {}
    for (row, col) in sorted(ws._cells):
        cell = ws._cells[(row, col)]
        cached = None
        if cell.data_type == "f" and ws_values is not None:
            cached = ws_values.cell(row=row, column=col).value
        record = _encode_cell(cell, cached, styles)
        if record is not None:
            cells[f"{get_column_letter(col)}{row}"] = record
    sheet["cells"] = cells
    return sheet


def _export_defined_names(names) -> List[Dict[str, Any]]:
    result = []
    for dn in names:
        data = {"name": dn.name, "value": dn.attr_text}
        data.update(pick_attrs(dn, ["comment", "hidden"]))
        if not data.get("hidden"):
            data.pop("hidden", None)
        result.append(data)
    return result


def _export_custom_properties(wb) -> List[Dict[str, Any]]:
    """User-defined document properties (File > Info > Properties > Custom)."""
    result = []
    for prop in wb.custom_doc_props.props:
        value = prop.value
        if isinstance(value, datetime.datetime):
            value = value.isoformat()
        result.append({"name": prop.name, "type": type(prop).__name__, "value": value})
    return result


def _export_chartsheets(wb) -> List[Dict[str, Any]]:
    result = []
    for position, sheet in enumerate(wb._sheets):
        if sheet not in wb.chartsheets:
            continue
        data: Dict[str, Any] = {"name": sheet.title, "position": position}
        if sheet.sheet_state != "visible":
            data["state"] = sheet.sheet_state
        data["charts"] = [chart_to_json(chart) for chart in sheet._charts]
        result.append(data)
    return result


def _export_external_links(wb) -> List[Dict[str, Any]]:
    """Links to other workbooks, in order: formulas refer to them as [1], [2]..."""
    result = []
    for link in getattr(wb, "_external_links", []):
        rel = link.file_link
        data: Dict[str, Any] = {"target": rel.Target}
        if rel.TargetMode:
            data["targetMode"] = rel.TargetMode
        data["type"] = rel.Type
        data["xml"] = tostring(link.to_tree()).decode("utf-8")
        result.append(data)
    return result


def _extract_vba(path: Path) -> Dict[str, bytes]:
    """Read VBA related parts (xl/vbaProject.bin etc.) directly from the archive."""
    parts = {}
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        for name in names:
            if name.startswith("xl/vba") or name.startswith("customUI"):
                parts[name] = archive.read(name)
        # Ribbon customizations are referenced from the package relationships.
        if any(name.startswith("customUI") for name in parts) and "_rels/.rels" in names:
            parts["_rels/.rels"] = archive.read("_rels/.rels")
    return parts


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def export_model(path: Union[str, Path], cached_values: bool = True) -> Dict[str, Any]:
    """Load an Excel file and convert it into the xlsx2txt model.

    The model is a plain dict with the keys ``manifest``, ``workbook``,
    ``styles``, ``sheets``, ``theme``, ``vba``, ``vbaSources``, ``media`` and ``pivots``.

    Args:
        path: Path to .xlsx / .xlsm file.
        cached_values: Also store the last calculated values of formula cells
            (read-only information, useful for review and diffs).
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    with zipfile.ZipFile(path) as archive:
        # Detect macros by content: git textconv passes temporary file names.
        is_macro = "xl/vbaProject.bin" in archive.namelist()
    # Pass the content, not the path: openpyxl rejects names without an Excel
    # extension, and git textconv may hand over temporary files.
    content = path.read_bytes()
    wb = load_workbook(BytesIO(content), data_only=False, keep_links=True, rich_text=True)
    wb_values = load_workbook(BytesIO(content), data_only=True) if cached_values else None

    styles = StyleTable()
    # Index 0 is the default style of the workbook.
    default_cell = OpenpyxlCell(wb.worksheets[0]) if wb.worksheets else None
    if default_cell is not None:
        styles.add(default_cell)

    images, warnings = extract_images(path)
    media: Dict[str, bytes] = {}
    pivots = export_pivots(wb.worksheets)

    sheets = []
    for ws in wb.worksheets:
        ws_values = wb_values[ws.title] if wb_values is not None else None
        sheet = _export_sheet(ws, ws_values, styles)
        sheet_images = []
        for content, extension, anchor in images.get(ws.title, []):
            name = media_name(content, extension)
            media[name] = content
            sheet_images.append({"file": name, "anchor": anchor_to_json(anchor)})
        if sheet_images:
            sheet["images"] = sheet_images
        if ws._charts:
            sheet["charts"] = [chart_to_json(chart) for chart in ws._charts]
        if ws.title in pivots["sheets"]:
            sheet["pivotTables"] = pivots["sheets"][ws.title]
        # Keep cells last: they are the largest part of the file.
        sheet["cells"] = sheet.pop("cells")
        sheets.append(sheet)

    workbook: Dict[str, Any] = {
        "epoch": 1904 if wb.epoch.year == 1904 else 1900,
        "sheets": [ws.title for ws in wb.worksheets],
        "activeSheet": wb.worksheets.index(wb.active) if wb.active in wb.worksheets else 0,
        "properties": pick_attrs(wb.properties, PROPERTY_ATTRS),
        "definedNames": _export_defined_names(wb.defined_names.values()),
    }
    custom = _export_custom_properties(wb)
    if custom:
        workbook["customProperties"] = custom
    chartsheets = _export_chartsheets(wb)
    if chartsheets:
        workbook["chartsheets"] = chartsheets
    external_links = _export_external_links(wb)
    if external_links:
        workbook["externalLinks"] = external_links
    calc = pick_attrs(wb.calculation, CALC_ATTRS)
    if calc:
        workbook["calculation"] = calc

    theme = wb.loaded_theme
    if isinstance(theme, bytes):
        theme = theme.decode("utf-8")

    vba = _extract_vba(path) if is_macro else {}
    vba_sources: Dict[str, str] = {}
    if vba:
        vba_sources, vba_warnings = extract_vba_sources(path)
        warnings.extend(vba_warnings)

    manifest = {
        "format": FORMAT_NAME,
        "formatVersion": FORMAT_VERSION,
        "generator": f"xlsx2txt {__version__}",
        "source": {"name": path.name, "type": path.suffix.lower().lstrip(".") or "xlsx"},
        "sheetCount": len(sheets),
        "cellCount": sum(len(s["cells"]) for s in sheets),
        "hasVba": bool(vba),
        "warnings": warnings,
    }

    return {
        "manifest": manifest,
        "workbook": workbook,
        "styles": styles.to_json(),
        "sheets": sheets,
        "theme": theme,
        "vba": vba,
        "vbaSources": vba_sources,
        "media": media,
        "pivots": pivots["files"],
    }
