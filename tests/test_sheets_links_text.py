"""Tests for chart sheets, external links and the text rendering (cat)."""

import subprocess
import sys

import pytest
from click.testing import CliRunner
from openpyxl import Workbook, load_workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.packaging.relationship import Relationship
from openpyxl.workbook.external_link.external import (
    ExternalBook,
    ExternalCell,
    ExternalLink,
    ExternalRow,
    ExternalSheetData,
    ExternalSheetDataSet,
    ExternalSheetNames,
)

from xlsx2txt import diff_models, export_model, export_xlsx, import_dir
from xlsx2txt.cli import main
from xlsx2txt.converter import roundtrip_diff
from xlsx2txt.importer import _vba_archive
from xlsx2txt.text import render_text

from tests.conftest import FIXTURES_DIR


@pytest.fixture
def linked_xlsx(tmp_path):
    path = tmp_path / "linked.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for i in range(1, 4):
        ws.append([f"k{i}", i])
    ws["C1"] = "=[1]Other!A1"

    pie = PieChart()
    pie.add_data(Reference(ws, min_col=2, min_row=1, max_row=3))
    chartsheet = wb.create_chartsheet("Pie", 0)
    chartsheet.add_chart(pie)
    wb.create_sheet("Last")

    book = ExternalBook(
        sheetNames=ExternalSheetNames(sheetName=["Other"]),
        sheetDataSet=ExternalSheetDataSet(sheetData=[
            ExternalSheetData(sheetId=0, row=[ExternalRow(r=1, cell=[ExternalCell(r="A1", v="42")])]),
        ]),
        id="rId1",
    )
    link = ExternalLink(externalBook=book)
    link.file_link = Relationship(type="externalLinkPath", Target="other.xlsx", TargetMode="External")
    wb._external_links.append(link)
    wb.active = 2  # "Last"
    wb.save(path)
    return path


def test_chartsheet_and_links_exported(linked_xlsx):
    model = export_model(linked_xlsx)
    workbook = model["workbook"]
    assert workbook["sheets"] == ["Data", "Last"]
    assert workbook["activeSheet"] == 1
    (chartsheet,) = workbook["chartsheets"]
    assert chartsheet["name"] == "Pie" and chartsheet["position"] == 0
    assert "pieChart" in chartsheet["charts"][0]["xml"]
    (link,) = workbook["externalLinks"]
    assert link["target"] == "other.xlsx"
    assert "Other" in link["xml"]
    assert model["manifest"]["warnings"] == []


def test_chartsheet_and_links_roundtrip(linked_xlsx, tmp_path):
    out_dir = tmp_path / "linked"
    model = export_xlsx(linked_xlsx, out_dir)
    assert roundtrip_diff(model) == []

    restored = tmp_path / "restored.xlsx"
    import_dir(out_dir, restored)
    wb = load_workbook(restored, keep_links=True)
    assert wb.sheetnames == ["Pie", "Data", "Last"]
    assert wb.active.title == "Last"
    assert wb._external_links[0].file_link.Target == "other.xlsx"
    assert wb["Data"]["C1"].value == "=[1]Other!A1"
    assert diff_models(model, export_model(restored), ignore_cached=True) == []


def test_render_text(linked_xlsx):
    text = render_text(export_model(linked_xlsx))
    assert "external link [1] -> other.xlsx" in text
    assert "=== Data ===\nA1: 'k1'\nB1: 1\nC1: =[1]Other!A1\n" in text
    assert "=== Pie (chart sheet) ===\nchart" in text


def test_render_text_styles_and_details():
    model = export_model(FIXTURES_DIR / "complex.xlsx")
    text = render_text(model)
    assert "A1: 'Sales Report Q1 2025'  [font 14pt bold; align horizontal=center]" in text
    assert "merged A1:D1" in text
    assert "D3: =B3*C3\n" in text
    assert "[font" not in render_text(model, styles=False)


def test_cat_command_accepts_file_and_dir(tmp_path):
    runner = CliRunner()
    from_file = runner.invoke(main, ["cat", str(FIXTURES_DIR / "complex.xlsx")])
    assert from_file.exit_code == 0, from_file.output
    out_dir = tmp_path / "complex"
    export_xlsx(FIXTURES_DIR / "complex.xlsx", out_dir)
    from_dir = runner.invoke(main, ["cat", str(out_dir)])
    assert from_dir.output == from_file.output


def test_macros_detected_by_content(tmp_path):
    source = tmp_path / "macro.xlsm"
    wb = Workbook()
    wb.vba_archive = _vba_archive({"xl/vbaProject.bin": b"dummy"})
    wb.save(source)
    renamed = tmp_path / "tmp_git_file"  # git textconv may pass files without extension
    renamed.write_bytes(source.read_bytes())
    assert export_model(renamed)["vba"] == {"xl/vbaProject.bin": b"dummy"}


def test_git_textconv(tmp_path):
    git = subprocess.run(["git", "--version"], capture_output=True)
    if git.returncode != 0:
        pytest.skip("git is not available")
    repo = tmp_path / "repo"
    repo.mkdir()

    def run(*args):
        return subprocess.run(args, cwd=repo, capture_output=True, text=True, check=True).stdout

    run("git", "init", "-q")
    run("git", "config", "user.email", "test@example.com")
    run("git", "config", "user.name", "Test")
    run("git", "config", "diff.xlsx.textconv", f'"{sys.executable}" -m xlsx2txt.cli cat')
    (repo / ".gitattributes").write_text("*.xlsx diff=xlsx\n")
    (repo / "report.xlsx").write_bytes((FIXTURES_DIR / "complex.xlsx").read_bytes())
    run("git", "add", ".")
    run("git", "commit", "-qm", "init")

    wb = load_workbook(repo / "report.xlsx")
    wb["Data"]["B3"] = 150
    wb.save(repo / "report.xlsx")

    diff = run("git", "--no-pager", "diff")
    assert "-B3: 100\n+B3: 150" in diff
