"""Worksheet and workbook elements openpyxl drops are kept (xmlElements)."""

import zipfile

from openpyxl import Workbook, load_workbook

from xlsx2txt import diff_models, export_model, export_xlsx, import_dir
from xlsx2txt.converter import roundtrip_diff
from xlsx2txt.elements import SHEET_ORDER, WORKBOOK_ORDER
from xlsx2txt.xmlfrag import child_spans, local_name, split_children

R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
X15 = "http://schemas.microsoft.com/office/spreadsheetml/2010/11/main"
X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"

SHEET_ELEMENTS = {
    "sheetCalcPr": '<sheetCalcPr fullCalcOnLoad="1"/>',
    "protectedRanges": '<protectedRanges><protectedRange sqref="A1:A3" name="Inputs" '
                       'algorithmName="SHA-512" hashValue="abc=" saltValue="def=" spinCount="100000"/>'
                       '</protectedRanges>',
    "sortState": '<sortState ref="A1:A3"><sortCondition descending="1" ref="A1:A3"/></sortState>',
    "dataConsolidate": '<dataConsolidate function="sum"/>',
    "customSheetViews": '<customSheetViews><customSheetView guid="{11111111-1111-1111-1111-111111111111}" '
                        'scale="80"><pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" '
                        'footer="0.3"/></customSheetView></customSheetViews>',
    "phoneticPr": '<phoneticPr fontId="1" type="noConversion"/>',
    "cellWatches": '<cellWatches><cellWatch r="A1"/></cellWatches>',
    "ignoredErrors": '<ignoredErrors><ignoredError sqref="A1:A3" numberStoredAsText="1"/></ignoredErrors>',
    "webPublishItems": '<webPublishItems count="1"><webPublishItem id="1" divId="d1" sourceType="sheet" '
                       'destinationFile="page.htm"/></webPublishItems>',
}
WORKBOOK_ELEMENTS = {
    "fileVersion": '<fileVersion appName="xl" lastEdited="7" lowestEdited="7" rupBuild="27425"/>',
    "fileSharing": '<fileSharing readOnlyRecommended="1" userName="Olena"/>',
    "customWorkbookViews": '<customWorkbookViews><customWorkbookView name="Review" '
                           'guid="{22222222-2222-2222-2222-222222222222}" windowWidth="800" windowHeight="600" '
                           'activeSheetId="1"/></customWorkbookViews>',
    "fileRecoveryPr": '<fileRecoveryPr repairLoad="1"/>',
}
WORKBOOK_EXT = (f'<ext uri="{{140A7094-0E35-4892-8432-C4D2E57EDEB5}}" xmlns:x15="{X15}">'
                '<x15:workbookPr chartTrackingRefBase="1"/></ext>')
SLICER_EXT = (f'<ext uri="{{BBE1A952-AA13-448e-AADC-164F8A28A991}}" xmlns:x14="{X14}"><x14:slicerCaches>'
              '<x14:slicerCache r:id="rId77"/></x14:slicerCaches></ext>')


def _reorder(xml, extra, order, closing):
    root, spans = child_spans(xml)
    children = {local_name(xml[a:b]): xml[a:b] for a, b in spans}
    children.update(extra)
    return root + "".join(children[n] for n in order if n in children) + closing


def _source(path, background=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in range(1, 4):
        ws.cell(row, 1, str(row))
    wb.create_sheet("Other")
    wb.save(path)
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    extra = dict(SHEET_ELEMENTS)
    if background:
        extra["picture"] = '<picture r:id="rId5"/>'
    sheet = parts["xl/worksheets/sheet1.xml"].decode().replace("<worksheet ", f'<worksheet xmlns:r="{R}" ', 1)
    parts["xl/worksheets/sheet1.xml"] = _reorder(sheet, extra, SHEET_ORDER, "</worksheet>").encode()
    workbook = parts["xl/workbook.xml"].decode()
    extra = dict(WORKBOOK_ELEMENTS, extLst=f"<extLst>{WORKBOOK_EXT}{SLICER_EXT}</extLst>")
    parts["xl/workbook.xml"] = _reorder(workbook, extra, WORKBOOK_ORDER, "</workbook>").encode()
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)


def _names(fragments):
    return [local_name(f) for f in fragments]


def test_elements_exported(tmp_path):
    source = tmp_path / "elements.xlsx"
    _source(source)
    model = export_model(source)
    assert _names(model["sheets"][0]["xmlElements"]) == list(SHEET_ELEMENTS)
    # fileVersion is rewritten by Excel on every save and is left out.
    assert _names(model["workbook"]["xmlElements"]) == ["fileSharing", "customWorkbookViews", "fileRecoveryPr"]
    # The workbook extension without relationships is kept, the slicer one is not.
    main = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    assert [e["xml"] for e in model["workbook"]["extensions"]] == [
        WORKBOOK_EXT.replace('main">', f'main" {main}>', 1)]
    assert "xmlElements" not in model["sheets"][1]
    assert model["manifest"]["warnings"] == []


def test_elements_roundtrip(tmp_path):
    source = tmp_path / "elements.xlsx"
    _source(source)
    out_dir = tmp_path / "out"
    model = export_xlsx(source, out_dir)
    assert roundtrip_diff(model) == []

    restored = tmp_path / "restored.xlsx"
    import_dir(out_dir, restored)
    with zipfile.ZipFile(restored) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml").decode()
        workbook = archive.read("xl/workbook.xml").decode()
    names = [local_name(c) for c in split_children(sheet)[1]]
    assert names == sorted(names, key=SHEET_ORDER.index)  # schema order
    assert set(SHEET_ELEMENTS) <= set(names)
    workbook_names = [local_name(c) for c in split_children(workbook)[1]]
    assert workbook_names == sorted(workbook_names, key=WORKBOOK_ORDER.index)
    assert {"fileSharing", "customWorkbookViews", "fileRecoveryPr", "extLst"} <= set(workbook_names)
    assert "<x15:workbookPr chartTrackingRefBase=\"1\"/>" in workbook and "slicerCache" not in workbook
    load_workbook(restored)
    assert diff_models(model, export_model(restored), ignore_cached=True) == []


def test_elements_in_diff_and_warnings(tmp_path):
    source = tmp_path / "elements.xlsx"
    _source(source, background=True)
    model = export_model(source)
    assert "Sheet 'Data': background picture are not exported" in model["manifest"]["warnings"]

    edited = export_model(source)
    edited["sheets"][0]["xmlElements"] = [x for x in edited["sheets"][0]["xmlElements"]
                                          if local_name(x) != "cellWatches"]
    edited["workbook"]["xmlElements"][0] = edited["workbook"]["xmlElements"][0].replace("Olena", "Taras")
    assert diff_models(model, edited) == [
        "workbook <fileSharing>: changed",
        "[Data] <cellWatches>: removed",
    ]
