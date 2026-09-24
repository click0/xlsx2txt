"""Tests for directory storage, diff and verification."""

import json

import pytest

from xlsx2txt import diff_models, export_xlsx, load_model, verify_dir
from xlsx2txt.compare import validate_model
from xlsx2txt.storage import FormatError, check_checksums, dumps, read_model, sheet_file_names

from tests.conftest import FIXTURES_DIR


def test_layout(tmp_path):
    out_dir = tmp_path / "out"
    export_xlsx(FIXTURES_DIR / "complex.xlsx", out_dir)
    for rel in [
        "manifest.json", "data/workbook.json", "data/sheets/_index.json", "data/sheets/Data.json",
        "data/styles/fonts.json", "data/styles/cellStyles.json", "_verify/checksums.json",
    ]:
        assert (out_dir / rel).is_file(), rel
    index = json.loads((out_dir / "data/sheets/_index.json").read_text())
    assert index == [{"name": "Data", "file": "Data.json"}]


def test_one_cell_per_line(tmp_path):
    out_dir = tmp_path / "out"
    export_xlsx(FIXTURES_DIR / "complex.xlsx", out_dir)
    text = (out_dir / "data/sheets/Data.json").read_text(encoding="utf-8")
    assert '    "D3": {"t": "f", "f": "=B3*C3"},\n' in text


def test_dumps_is_valid_json():
    obj = {"a": [1, 2, {"b": "x" * 200}], "c": {"d": {"e": None}}, "long": list(range(100))}
    assert json.loads(dumps(obj)) == obj


def test_sheet_file_names():
    names = sheet_file_names(["Sheet", "a/b", "A<B", "a>b", "CON", "_index", ""])
    assert names == ["Sheet.json", "a_b.json", "A_B_2.json", "a_b_3.json", "_CON.json", "_index_2.json", "sheet_2.json"]


def test_verify_ok(tmp_path):
    out_dir = tmp_path / "out"
    export_xlsx(FIXTURES_DIR / "complex.xlsx", out_dir)
    report = verify_dir(out_dir, against=FIXTURES_DIR / "complex.xlsx")
    assert report["ok"], report


def test_verify_detects_modification(tmp_path):
    out_dir = tmp_path / "out"
    export_xlsx(FIXTURES_DIR / "simple.xlsx", out_dir)
    sheet = out_dir / "data/sheets/Sheet.json"
    sheet.write_text(sheet.read_text().replace("Hello", "Bye"))

    checks = check_checksums(out_dir)
    assert checks["modified"] == ["data/sheets/Sheet.json"]
    report = verify_dir(out_dir, against=FIXTURES_DIR / "simple.xlsx")
    assert not report["ok"]
    assert any("Hello" in line for line in report["against"])


def test_edited_directory_is_imported(tmp_path):
    from xlsx2txt import import_dir
    from openpyxl import load_workbook

    out_dir = tmp_path / "out"
    export_xlsx(FIXTURES_DIR / "simple.xlsx", out_dir)
    sheet = out_dir / "data/sheets/Sheet.json"
    data = json.loads(sheet.read_text())
    data["cells"]["B2"] = {"v": 7, "t": "n"}
    sheet.write_text(json.dumps(data))

    import_dir(out_dir, tmp_path / "edited.xlsx")
    ws = load_workbook(tmp_path / "edited.xlsx").active
    assert ws["A1"].value == "Hello" and ws["B2"].value == 7


def test_diff_reports_changes(tmp_path):
    a = load_model(FIXTURES_DIR / "complex.xlsx")
    b = load_model(FIXTURES_DIR / "complex.xlsx")
    cells = b["sheets"][0]["cells"]
    cells["B3"]["v"] = 101
    del cells["A5"]
    cells["F1"] = {"v": "new", "t": "s"}
    cells["A3"]["s"] = cells["A2"]["s"]  # bold

    changes = diff_models(a, b)
    assert "[Data] B3.v: 100 -> 101" in changes
    assert "[Data] A5: removed (was 'Gizmo')" in changes
    assert "[Data] F1: added 'new'" in changes
    assert any(line.startswith("[Data] A3.style.font") for line in changes)


def test_diff_ignores_style_index_renumbering():
    a = load_model(FIXTURES_DIR / "styled.xlsx")
    b = load_model(FIXTURES_DIR / "styled.xlsx")
    styles = b["styles"]["cellStyles"]
    styles.append(styles[1])
    b["sheets"][0]["cells"]["A1"]["s"] = len(styles) - 1
    assert diff_models(a, b) == []


def test_validate_model_errors():
    model = load_model(FIXTURES_DIR / "simple.xlsx")
    model["sheets"][0]["cells"]["A1"]["s"] = 99
    model["sheets"][0]["cells"]["A2"] = {"t": "x"}
    errors = validate_model(model)
    assert any("invalid style index" in e for e in errors)
    assert any("unknown type" in e for e in errors)


def test_read_model_rejects_foreign_dir(tmp_path):
    with pytest.raises(FormatError):
        read_model(tmp_path)


def test_export_refuses_foreign_non_empty_dir(tmp_path):
    (tmp_path / "keep.txt").write_text("x")
    with pytest.raises(FileExistsError):
        export_xlsx(FIXTURES_DIR / "simple.xlsx", tmp_path)
    assert (tmp_path / "keep.txt").exists()


def test_reexport_removes_stale_sheet_files(tmp_path):
    out_dir = tmp_path / "out"
    export_xlsx(FIXTURES_DIR / "complex.xlsx", out_dir)
    export_xlsx(FIXTURES_DIR / "simple.xlsx", out_dir)
    assert not (out_dir / "data/sheets/Data.json").exists()
    assert verify_dir(out_dir)["ok"]
