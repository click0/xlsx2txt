"""Shapes (text boxes, arrows, groups) are exported and restored."""

import re
import zipfile
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from PIL import Image as PILImage

from xlsx2txt import diff_models, export_model, export_xlsx, import_dir
from xlsx2txt.converter import roundtrip_diff
from xlsx2txt.shapes import drawing_children, shapes_from_drawing
from xlsx2txt.storage import write_model
from xlsx2txt.text import render_text

XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
HYPERLINK = f"{R}/hyperlink"
IMAGE = f"{R}/image"


def _anchor(col, row, body):
    return (f"<xdr:twoCellAnchor><xdr:from><xdr:col>{col}</xdr:col><xdr:colOff>0</xdr:colOff>"
            f"<xdr:row>{row}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>"
            f"<xdr:to><xdr:col>{col + 2}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{row + 2}</xdr:row>"
            f"<xdr:rowOff>0</xdr:rowOff></xdr:to>{body}<xdr:clientData/></xdr:twoCellAnchor>")


def _sp(shape_id, name, text, extra="", fill=""):
    return (f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{shape_id}" name="{name}">{extra}'
            f'</xdr:cNvPr><xdr:cNvSpPr txBox="1"/></xdr:nvSpPr><xdr:spPr><a:prstGeom prst="rect"><a:avLst/>'
            f'</a:prstGeom>{fill}</xdr:spPr><xdr:txBody><a:bodyPr/><a:lstStyle/>'
            f'<a:p><a:r><a:t>{text}</a:t></a:r></a:p></xdr:txBody></xdr:sp>')


# As Excel writes it: namespaces declared on the root only.
BELOW = _anchor(0, 2, _sp(9, "Under picture", "Below"))
SHAPES = [
    _anchor(3, 1, _sp(1, "TextBox 1", "Hello &amp; welcome")),
    f'<mc:AlternateContent xmlns:mc="{MC}"><mc:Choice xmlns:a14="{A14}" Requires="a14">'
    + _anchor(3, 5, _sp(2, "Rectangle 2", "Choice")) + "</mc:Choice><mc:Fallback/></mc:AlternateContent>",
    _anchor(3, 9, '<xdr:cxnSp macro=""><xdr:nvCxnSpPr><xdr:cNvPr id="3" name="Arrow 3"/><xdr:cNvCxnSpPr>'
                  '<a:stCxn id="1" idx="2"/><a:endCxn id="2" idx="0"/></xdr:cNvCxnSpPr></xdr:nvCxnSpPr>'
                  '<xdr:spPr><a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom></xdr:spPr></xdr:cxnSp>'),
    _anchor(8, 1, _sp(4, "Link 4", "Click", '<a:hlinkClick r:id="rId9"/>')),
    _anchor(8, 5, _sp(5, "Filled 5", "", fill='<a:blipFill><a:blip r:embed="rId10"/><a:stretch><a:fillRect/>'
                                              '</a:stretch></a:blipFill>')),
]
CONTROL = _anchor(8, 9, _sp(6, "Button 6", "Run", '<a:extLst><a:ext uri="{63B3BB69-23CF-44E3-9099-C40C66FF867C}">'
                                                  f'<a14:compatExt xmlns:a14="{A14}" spid="_x0000_s1025"/>'
                                                  '</a:ext></a:extLst>'))
NAMES = ["Under picture", "TextBox 1", "Rectangle 2", "Arrow 3", "Link 4", "Filled 5"]


def _png(color="red"):
    buffer = BytesIO()
    PILImage.new("RGB", (4, 4), color).save(buffer, format="PNG")
    return buffer.getvalue()


FILL = _png("blue")


def _excel_like(path):
    """A picture (id 1 after openpyxl) between shapes whose ids clash with it."""
    wb = Workbook()
    wb.active.title = "Data"
    wb.active["A1"] = "value"
    wb.active.add_image(Image(BytesIO(_png())), "A3")
    wb.create_sheet("Empty")
    wb.save(path)
    with zipfile.ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    drawing = parts["xl/drawings/drawing1.xml"].decode()
    root, children = drawing_children(drawing)
    new_root = (f'<xdr:wsDr xmlns:xdr="{XDR}" xmlns:a="{A}" xmlns:r="{R}">')
    # Re-prefix openpyxl's picture anchor so the drawing looks like Excel's.
    picture = re.sub(r"<(/?)(?!a:)(\w+)", r"<\1xdr:\2", children[0])
    picture = re.sub(r"<(/?)xdr:(blip|stretch|fillRect|prstGeom|avLst|off|ext|xfrm|picLocks)\b", r"<\1a:\2",
                     picture)
    parts["xl/drawings/drawing1.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + new_root + BELOW + picture
        + "".join(SHAPES) + CONTROL + "</xdr:wsDr>"
    ).encode()
    rels = parts["xl/drawings/_rels/drawing1.xml.rels"].decode()
    parts["xl/drawings/_rels/drawing1.xml.rels"] = rels.replace("</Relationships>", (
        f'<Relationship Id="rId9" Type="{HYPERLINK}" Target="https://example.com/?a=1&amp;b=2" '
        'TargetMode="External"/>'
        f'<Relationship Id="rId10" Type="{IMAGE}" Target="../media/fill.png"/></Relationships>'
    )).encode()
    parts["xl/media/fill.png"] = FILL
    with zipfile.ZipFile(path, "w") as target:
        for name, content in parts.items():
            target.writestr(name, content)


def test_drawing_children_and_classification():
    xml = f'<xdr:wsDr xmlns:xdr="{XDR}" xmlns:a="{A}" xmlns:r="{R}">' + "".join(SHAPES[:3]) + CONTROL + "</xdr:wsDr>"
    root, children = drawing_children(xml)
    assert root.startswith("<xdr:wsDr") and len(children) == 4
    shapes, unsupported, media = shapes_from_drawing(xml)
    assert unsupported == 1 and media == {}
    assert [s["name"] for s in shapes] == ["TextBox 1", "Rectangle 2", "Arrow 3"]
    assert shapes[0]["text"] == "Hello & welcome"
    assert "text" not in shapes[2]
    # Every fragment declares the namespaces it uses.
    assert shapes[0]["xml"].startswith(f'<xdr:twoCellAnchor xmlns:xdr="{XDR}" xmlns:a="{A}">')
    assert f'xmlns:mc="{MC}"' in shapes[1]["xml"] and f'xmlns:xdr="{XDR}"' in shapes[1]["xml"]
    # A reference to a relationship the drawing does not have cannot be kept.
    shapes, unsupported, _ = shapes_from_drawing(xml.replace("</xdr:wsDr>", SHAPES[3] + "</xdr:wsDr>"))
    assert unsupported == 2 and len(shapes) == 3


def test_shapes_exported(tmp_path):
    source = tmp_path / "shapes.xlsx"
    _excel_like(source)
    model = export_model(source)
    shapes = model["sheets"][0]["shapes"]
    assert [s["name"] for s in shapes] == NAMES
    assert shapes[0]["layer"] == 0  # drawn below the picture
    assert all("layer" not in s for s in shapes[1:])
    assert shapes[4]["rels"] == {"rId9": {"type": "hyperlink", "target": "https://example.com/?a=1&b=2",
                                          "external": True}}
    fill = shapes[5]["rels"]["rId10"]
    assert fill["type"] == "image" and model["media"][fill["file"]] == FILL
    assert len(model["media"]) == 2  # the picture and the fill
    assert model["manifest"]["warnings"] == [
        "Sheet 'Data': 1 form control(s) or shape(s) with unsupported links are not exported"
    ]


def test_shapes_roundtrip(tmp_path):
    source = tmp_path / "shapes.xlsx"
    _excel_like(source)
    out_dir = tmp_path / "out"
    model = export_xlsx(source, out_dir)
    assert roundtrip_diff(model) == []

    restored = tmp_path / "restored.xlsx"
    import_dir(out_dir, restored)
    with zipfile.ZipFile(restored) as archive:
        drawing = archive.read("xl/drawings/drawing1.xml").decode()
        rels = archive.read("xl/drawings/_rels/drawing1.xml.rels").decode()
        types = archive.read("[Content_Types].xml").decode()
        ids = re.findall(r'cNvPr id="(\d+)"', drawing)
        assert len(ids) == len(set(ids)) == 7  # picture + 6 shapes, clash with the picture resolved
        # The connector still points at the renumbered shapes.
        start, end = re.search(r'stCxn id="(\d+)".*?endCxn id="(\d+)"', drawing).groups()
        names = dict(re.findall(r'cNvPr id="(\d+)" name="([^"]+)"', drawing))
        assert (names[start], names[end]) == ("TextBox 1", "Rectangle 2")
        # Z-order: the first shape is still below the picture.
        order = [name for _, name in re.findall(r'cNvPr id="(\d+)" name="([^"]+)"', drawing)]
        assert order[0] == "Under picture" and order[1] != "TextBox 1"
        # Relationships of the shapes are recreated.
        link_id = re.search(r'hlinkClick r:id="([^"]+)"', drawing).group(1)
        assert re.search(rf'Id="{link_id}"[^>]*Target="https://example.com/\?a=1&amp;b=2" TargetMode="External"',
                         rels)
        fill_id = re.search(r'blipFill><a:blip r:embed="([^"]+)"', drawing).group(1)
        fill_target = re.search(rf'Id="{fill_id}"[^>]*Target="([^"]+)"', rels).group(1)
        assert archive.read("xl/" + fill_target.replace("../", "")) == FILL
        assert 'Extension="png"' in types
    load_workbook(restored)  # still readable
    again = export_model(restored)
    assert diff_models(model, again, ignore_cached=True) == []
    assert [s["name"] for s in again["sheets"][0]["shapes"]] == NAMES


def test_shapes_on_sheet_without_drawing(tmp_path):
    source = tmp_path / "shapes.xlsx"
    _excel_like(source)
    out_dir = tmp_path / "out"
    model = export_xlsx(source, out_dir)
    # Move the shapes to the sheet that has no drawing yet.
    model["sheets"][1]["shapes"] = model["sheets"][0].pop("shapes")
    write_model(model, out_dir)
    restored = tmp_path / "restored.xlsx"
    import_dir(out_dir, restored)
    with zipfile.ZipFile(restored) as archive:
        sheet = archive.read("xl/worksheets/sheet2.xml").decode()
        types = archive.read("[Content_Types].xml").decode()
        assert re.search(r'<drawing r:id="rId\d+"/></worksheet>', sheet)
        assert types.count("drawing+xml") == 2
    again = export_model(restored)
    assert "shapes" not in again["sheets"][0]
    shapes = again["sheets"][1]["shapes"]
    assert [s["name"] for s in shapes] == NAMES
    assert all("layer" not in s for s in shapes)  # no pictures there
    assert again["media"] == model["media"]


def test_shape_changes_in_diff_and_text(tmp_path):
    source = tmp_path / "shapes.xlsx"
    _excel_like(source)
    model = export_model(source)
    text = render_text(model)
    assert "shape 'TextBox 1' at D2: 'Hello & welcome'" in text
    assert "shape 'Arrow 3' at D10" in text
    assert "shape 'Link 4' at I2: 'Click'  <link https://example.com/?a=1&b=2>" in text

    edited = export_model(source)
    shapes = edited["sheets"][0]["shapes"]
    shapes[1]["xml"] = shapes[1]["xml"].replace("Hello &amp; welcome", "Bye")
    shapes[1]["text"] = "Bye"
    shapes[4]["rels"]["rId9"]["target"] = "https://example.org/"
    del shapes[0]["layer"]
    del shapes[3]
    assert diff_models(model, edited) == [
        "[Data] shape 'Under picture': changed",
        "[Data] shape 'TextBox 1': text changed",
        "[Data] shape 'Arrow 3': removed",
        "[Data] shape 'Link 4': changed",
    ]
