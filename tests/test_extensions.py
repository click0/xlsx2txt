"""Worksheet extensions (sparklines, x14 rules) and threaded comments."""

import re
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import DataBarRule

from xlsx2txt import diff_models, export_model, export_xlsx, import_dir
from xlsx2txt.compare import validate_model
from xlsx2txt.converter import roundtrip_diff
from xlsx2txt.text import render_text

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
XM = "http://schemas.microsoft.com/office/excel/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
X14AC = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac"
TC_NS = "http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments"
BAR_ID = "{00000000-0000-0000-0000-00000000AB01}"
THREAD_ID = "{00000000-0000-0000-0000-0000000000C1}"
REPLY_ID = "{00000000-0000-0000-0000-0000000000C2}"
ANNA = "{00000000-0000-0000-0000-0000000000A1}"
BOHDAN = "{00000000-0000-0000-0000-0000000000B1}"

# As Excel writes them: xm is declared on the worksheet root here.
CF_EXT = (f'<ext uri="{{78C0D931-6437-407d-A8EE-F0AAD7539E65}}" xmlns:x14="{X14}"><x14:conditionalFormattings>'
          f'<x14:conditionalFormatting><x14:cfRule type="dataBar" id="{BAR_ID}"><x14:dataBar minLength="0" '
          'maxLength="100" negativeBarColorSameAsPositive="0" axisPosition="middle"><x14:cfvo type="autoMin"/>'
          '<x14:cfvo type="autoMax"/><x14:negativeFillColor rgb="FFFF0000"/></x14:dataBar></x14:cfRule>'
          '<xm:sqref>B1:B4</xm:sqref></x14:conditionalFormatting></x14:conditionalFormattings></ext>')
DV_EXT = (f'<ext uri="{{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}}" xmlns:x14="{X14}"><x14:dataValidations count="1">'
          '<x14:dataValidation type="list" allowBlank="1"><x14:formula1><xm:f>Lists!$A$1:$A$2</xm:f>'
          '</x14:formula1><xm:sqref>C1:C4</xm:sqref></x14:dataValidation></x14:dataValidations></ext>')
SPARK_EXT = (f'<ext uri="{{05C60535-1F16-4fd2-B633-F4F36F0B64E0}}" xmlns:x14="{X14}"><x14:sparklineGroups>'
             '<x14:sparklineGroup displayEmptyCellsAs="gap"><x14:colorSeries rgb="FF376092"/><x14:sparklines>'
             '<x14:sparkline><xm:f>Data!B1:B4</xm:f><xm:sqref>D1</xm:sqref></x14:sparkline></x14:sparklines>'
             '</x14:sparklineGroup></x14:sparklineGroups></ext>')
OTHER_EXT = '<ext uri="{00000000-1111-2222-3333-444444444444}"><x14ac:note xmlns:x14ac="' + X14AC + '"/></ext>'
SLICER_EXT = (f'<ext uri="{{A8765BA9-456A-4dab-B4F3-ACF838C121DE}}" xmlns:x14="{X14}"><x14:slicerList>'
              '<x14:slicer r:id="rId7"/></x14:slicerList></ext>')

THREADS_XML = (
    f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<ThreadedComments xmlns="{TC_NS}" '
    f'xmlns:x="{MAIN}"><threadedComment ref="A2" dT="2026-03-01T10:00:00.00" personId="{ANNA}" '
    f'id="{THREAD_ID}"><text>@Bohdan check &amp; confirm</text><mentions><mention mentionpersonId="{BOHDAN}" '
    'mentionId="{00000000-0000-0000-0000-0000000000D1}" startIndex="0" length="7"/></mentions></threadedComment>'
    f'<threadedComment ref="A2" dT="2026-03-02T08:00:00.00" personId="{BOHDAN}" id="{REPLY_ID}" '
    f'parentId="{THREAD_ID}" done="1"><text>Done.\nAll good.</text></threadedComment></ThreadedComments>'
)
PERSONS_XML = (
    f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<personList xmlns="{TC_NS}" xmlns:x="{MAIN}">'
    f'<person displayName="Anna" id="{ANNA}" userId="anna@example.com" providerId="None"/>'
    f'<person displayName="Bohdan" id="{BOHDAN}" userId="S::bohdan@example.com::1" providerId="AD"/>'
    '</personList>'
)


def _excel_like(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row, value in enumerate([5, -3, 8, 2], start=1):
        ws.cell(row, 2, value)
    ws.conditional_formatting.add("B1:B4", DataBarRule(start_type="min", end_type="max", color="FF638EC6"))
    ws["A2"] = "Total"
    ws["A2"].comment = Comment("[Threaded comment]\n\nComment:\n    @Bohdan check & confirm", f"tc={THREAD_ID}")
    wb.create_sheet("Lists").append(["Yes"])
    wb.save(path)

    with zipfile.ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    sheet = sheet.replace(
        "<worksheet ", f'<worksheet xmlns:mc="{MC}" mc:Ignorable="x14ac" xmlns:x14ac="{X14AC}" xmlns:xm="{XM}" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
    link = (f'<extLst><ext uri="{{B025F937-C7B1-47D3-B67F-A62EFF666E3E}}" xmlns:x14="{X14}">'
            f'<x14:id>{BAR_ID}</x14:id></ext></extLst></cfRule>')
    sheet = re.sub(r"(<cfRule [^>]*>.*?)</cfRule>", lambda m: m.group(1) + link, sheet, count=1, flags=re.S)
    sheet = sheet.replace("</worksheet>",
                          f"<extLst>{CF_EXT}{DV_EXT}{SPARK_EXT}{OTHER_EXT}{SLICER_EXT}</extLst></worksheet>")
    parts["xl/worksheets/sheet1.xml"] = sheet.encode()

    rels = parts["xl/worksheets/_rels/sheet1.xml.rels"].decode()
    parts["xl/worksheets/_rels/sheet1.xml.rels"] = rels.replace("</Relationships>", (
        '<Relationship Id="rId9" Type="http://schemas.microsoft.com/office/2017/10/relationships/threadedComment" '
        'Target="../threadedComments/threadedComment1.xml"/></Relationships>')).encode()
    wb_rels = parts["xl/_rels/workbook.xml.rels"].decode()
    parts["xl/_rels/workbook.xml.rels"] = wb_rels.replace("</Relationships>", (
        '<Relationship Id="rId99" Type="http://schemas.microsoft.com/office/2017/10/relationships/person" '
        'Target="persons/person.xml"/></Relationships>')).encode()
    parts["xl/threadedComments/threadedComment1.xml"] = THREADS_XML.encode()
    parts["xl/persons/person.xml"] = PERSONS_XML.encode()
    types = parts["[Content_Types].xml"].decode()
    parts["[Content_Types].xml"] = types.replace("</Types>", (
        '<Override PartName="/xl/threadedComments/threadedComment1.xml" '
        'ContentType="application/vnd.ms-excel.threadedcomments+xml"/>'
        '<Override PartName="/xl/persons/person.xml" ContentType="application/vnd.ms-excel.person+xml"/>'
        '</Types>')).encode()
    with zipfile.ZipFile(path, "w") as target:
        for name, content in parts.items():
            target.writestr(name, content)


def test_extensions_exported(tmp_path):
    source = tmp_path / "ext.xlsx"
    _excel_like(source)
    model = export_model(source)
    sheet = model["sheets"][0]
    assert [e["type"] for e in sheet["extensions"]] == [
        "conditional formatting", "data validation", "sparklines", "extension {00000000-1111-2222-3333-444444444444}",
    ]
    # Namespaces inherited from the worksheet root are declared on each fragment.
    assert f'xmlns:xm="{XM}"' in sheet["extensions"][2]["xml"]
    assert f'xmlns:xm="{XM}"' not in sheet["extensions"][3]["xml"]
    assert sheet["conditionalFormatting"][0]["rules"][0]["extId"] == BAR_ID
    assert model["manifest"]["warnings"] == ["Sheet 'Data': slicers are not exported"]


def test_extensions_roundtrip(tmp_path):
    source = tmp_path / "ext.xlsx"
    _excel_like(source)
    out_dir = tmp_path / "out"
    model = export_xlsx(source, out_dir)
    assert roundtrip_diff(model) == []

    restored = tmp_path / "restored.xlsx"
    import_dir(out_dir, restored)
    with zipfile.ZipFile(restored) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml").decode()
    rule = re.search(r"<cfRule\b.*?</cfRule>", sheet, re.S).group(0)
    assert rule.endswith(f"<x14:id>{BAR_ID}</x14:id></ext></extLst></cfRule>")
    assert sheet.rstrip().endswith("</extLst></worksheet>")
    for ext in model["sheets"][0]["extensions"]:
        assert ext["xml"] in sheet
    load_workbook(restored)
    assert diff_models(model, export_model(restored), ignore_cached=True) == []


def test_threaded_comments(tmp_path):
    source = tmp_path / "ext.xlsx"
    _excel_like(source)
    out_dir = tmp_path / "out"
    model = export_xlsx(source, out_dir)
    records = model["sheets"][0]["threadedComments"]
    assert [r["id"] for r in records] == [THREAD_ID, REPLY_ID]
    assert records[0]["text"] == "@Bohdan check & confirm"
    assert "<mention " in records[0]["xml"]
    assert records[1] == {"ref": "A2", "id": REPLY_ID, "parentId": THREAD_ID, "personId": BOHDAN,
                          "dT": "2026-03-02T08:00:00.00", "done": "1", "text": "Done.\nAll good."}
    assert [p["displayName"] for p in model["workbook"]["persons"]] == ["Anna", "Bohdan"]
    assert validate_model(model) == []

    text = render_text(model)
    assert "thread at A2: Anna: '@Bohdan check & confirm'" in text
    assert "  reply [resolved]: Bohdan: 'Done.\\nAll good.'" in text

    restored = tmp_path / "restored.xlsx"
    import_dir(out_dir, restored)
    with zipfile.ZipFile(restored) as archive:
        names = archive.namelist()
        assert "xl/persons/person.xml" in names
        types = archive.read("[Content_Types].xml").decode()
        assert "threadedcomments+xml" in types and "person+xml" in types
    again = export_model(restored)
    assert again["sheets"][0]["threadedComments"] == records
    assert again["workbook"]["persons"] == model["workbook"]["persons"]
    # The legacy note Excel shows in older versions is still there.
    assert load_workbook(restored)["Data"]["A2"].comment.author == f"tc={THREAD_ID}"


def test_thread_changes_in_diff(tmp_path):
    source = tmp_path / "ext.xlsx"
    _excel_like(source)
    model = export_model(source)
    edited = export_model(source)
    edited["sheets"][0]["threadedComments"][0]["text"] = "Changed"
    del edited["sheets"][0]["threadedComments"][1]
    edited["sheets"][0]["extensions"].pop()
    assert diff_models(model, edited) == [
        f"[Data] comment at A2 {THREAD_ID}: '@Bohdan check & confirm' -> 'Changed'",
        f"[Data] comment at A2 {REPLY_ID}: removed",
        "[Data] extension 'extension {00000000-1111-2222-3333-444444444444}': removed",
    ]
    edited["sheets"][0]["threadedComments"][0]["personId"] = "{unknown}"
    assert any("unknown person" in e for e in validate_model(edited))


METADATA_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    f'<metadata xmlns="{MAIN}" xmlns:xda="http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray">'
    '<metadataTypes count="1"><metadataType name="XLDAPR" minSupportedVersion="120000" copy="1" pasteAll="1" '
    'pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" clearComments="1" assign="1" '
    'coerce="1" cellMeta="1"/></metadataTypes><futureMetadata name="XLDAPR" count="1"><bk><extLst>'
    '<ext uri="{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}"><xda:dynamicArrayProperties fDynamic="1" fCollapsed="0"/>'
    '</ext></extLst></bk></futureMetadata><cellMetadata count="1"><bk><rc t="1" v="0"/></bk></cellMetadata>'
    '</metadata>'
)
IGNORED = '<ignoredErrors><ignoredError sqref="A5" numberStoredAsText="1"/></ignoredErrors>'


def _dynamic_array(path, metadata=METADATA_XML):
    """A dynamic array formula as Excel 365 writes it, and a disabled error check."""
    from openpyxl.worksheet.formula import ArrayFormula

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row, value in enumerate([3, 1, 2], start=1):
        ws.cell(row, 1, value)
    ws["B1"] = ArrayFormula("B1:B3", "=_xlfn._xlws.SORT(A1:A3)")
    ws["A5"] = "007"
    ws.conditional_formatting.add("A1:A3", DataBarRule(start_type="min", end_type="max", color="FF638EC6"))
    wb.save(path)
    with zipfile.ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    sheet = sheet.replace('<c r="B1"', '<c r="B1" cm="1"', 1)
    sheet = sheet.replace("</worksheet>", IGNORED + "</worksheet>")
    parts["xl/worksheets/sheet1.xml"] = sheet.encode()
    parts["xl/metadata.xml"] = metadata.encode()
    wb_rels = parts["xl/_rels/workbook.xml.rels"].decode()
    parts["xl/_rels/workbook.xml.rels"] = wb_rels.replace("</Relationships>", (
        '<Relationship Id="rId98" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
        'sheetMetadata" Target="metadata.xml"/></Relationships>')).encode()
    types = parts["[Content_Types].xml"].decode()
    parts["[Content_Types].xml"] = types.replace("</Types>", (
        '<Override PartName="/xl/metadata.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/></Types>'
    )).encode()
    with zipfile.ZipFile(path, "w") as target:
        for name, content in parts.items():
            target.writestr(name, content)


def test_dynamic_arrays_and_ignored_errors(tmp_path):
    source = tmp_path / "dyn.xlsx"
    _dynamic_array(source)
    out_dir = tmp_path / "out"
    model = export_xlsx(source, out_dir)
    cell = model["sheets"][0]["cells"]["B1"]
    assert cell["cm"] == 1 and cell["fType"] == "array"
    assert model["metadata"] == METADATA_XML
    assert (out_dir / "data" / "metadata.xml").read_text(encoding="utf-8") == METADATA_XML
    assert model["sheets"][0]["ignoredErrors"] == (
        f'<ignoredErrors xmlns="{MAIN}"><ignoredError sqref="A5" numberStoredAsText="1"/></ignoredErrors>')
    assert model["manifest"]["warnings"] == []
    assert "B1: =_xlfn._xlws.SORT(A1:A3) (spills over B1:B3)" in render_text(model)
    assert roundtrip_diff(model) == []

    # Move a shape onto the sheet too: <drawing> and <ignoredErrors> must both
    # land in schema order, around the rule's own <extLst>.
    model["sheets"][0]["conditionalFormatting"][0]["rules"][0]["extId"] = "{00000000-0000-0000-0000-0000000000E1}"
    model["sheets"][0]["shapes"] = [{"name": "Box", "xml": (
        '<xdr:twoCellAnchor xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><xdr:from><xdr:col>3</xdr:col>'
        '<xdr:colOff>0</xdr:colOff><xdr:row>1</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to>'
        '<xdr:col>5</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>3</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
        '<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="2" name="Box"/><xdr:cNvSpPr/></xdr:nvSpPr>'
        '<xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:sp><xdr:clientData/>'
        '</xdr:twoCellAnchor>')}]
    restored = tmp_path / "restored.xlsx"
    from xlsx2txt.importer import import_model
    import_model(model, restored)
    with zipfile.ZipFile(restored) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml").decode()
        assert archive.read("xl/metadata.xml").decode() == METADATA_XML
        assert "sheetMetadata" in archive.read("xl/_rels/workbook.xml.rels").decode()
        assert "sheetMetadata+xml" in archive.read("[Content_Types].xml").decode()
    assert re.search(r'<c r="B1" cm="1">', sheet)
    from xlsx2txt.xmlfrag import local_name, split_children
    order = [local_name(child) for child in split_children(sheet)[1]]
    assert order.index("pageMargins") < order.index("ignoredErrors") < order.index("drawing")
    assert "extLst" not in order  # the rule's own extLst stays inside <cfRule>
    assert "<x14:id>" in re.search(r"<cfRule\b.*?</cfRule>", sheet, re.S).group(0)
    again = export_model(restored)
    assert again["sheets"][0]["cells"]["B1"]["cm"] == 1
    assert again["sheets"][0]["ignoredErrors"] == model["sheets"][0]["ignoredErrors"]
    assert [shape["name"] for shape in again["sheets"][0]["shapes"]] == ["Box"]


def test_metadata_of_pictures_in_cells_is_not_kept(tmp_path):
    source = tmp_path / "rich.xlsx"
    _dynamic_array(source, METADATA_XML.replace("</metadata>", '<valueMetadata count="1"><bk><rc t="1" v="0"/>'
                                                                '</bk></valueMetadata></metadata>'))
    model = export_model(source)
    assert model["metadata"] is None
    assert "cm" not in model["sheets"][0]["cells"]["B1"]
    assert "cell metadata of pictures in cells / rich data types: 1 part(s) are not exported" in (
        model["manifest"]["warnings"])
