"""Round-trip fidelity checked independently of xlsx2txt's own model.

`verify` compares exports, so something the exporter never reads (e.g. page
headers before they were supported) cannot show up there. These tests open
each example and its restored copy with openpyxl and compare what openpyxl
sees: every cell's value, type and formatting, plus sheet and workbook
settings.
"""

import warnings
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from xlsx2txt import export_model
from xlsx2txt.importer import import_model

EXAMPLES = sorted((Path(__file__).resolve().parent.parent / "examples").glob("*.xlsx"))


def _side(side):
    return None if side is None or (side.style is None and side.color is None) else repr(side)


def _border(border):
    # An empty <left/> and a missing one mean the same: no line.
    sides = tuple(_side(getattr(border, name)) for name in ("left", "right", "top", "bottom", "diagonal"))
    return sides + (bool(border.diagonalUp), bool(border.diagonalDown))


def _value(value):
    if isinstance(value, ArrayFormula):
        return ("array", value.ref, value.text)
    return (type(value).__name__, str(value))


def _cell(cell):
    return {
        "value": _value(cell.value),
        "type": cell.data_type,
        "numFmt": cell.number_format,
        "font": repr(cell.font),
        "fill": repr(cell.fill),
        "border": _border(cell.border),
        "alignment": repr(cell.alignment),
        "protection": repr(cell.protection),
        "style": cell.style,
        "comment": (cell.comment.text, cell.comment.author) if cell.comment else None,
        "link": (cell.hyperlink.target, cell.hyperlink.location, cell.hyperlink.tooltip) if cell.hyperlink else None,
    }


def _dxf(wb, dxf_id):
    if dxf_id is None:
        return None
    styles = wb._differential_styles.styles
    return repr(styles[dxf_id]) if dxf_id < len(styles) else f"missing {dxf_id}"


def _resolve_dxfs(wb, xml):
    """Replace dxf ids in XML by the formats they point at."""
    import re
    return re.sub(r'(\w*[dD]xfId)="(\d+)"', lambda m: f'{m.group(1)}="{_dxf(wb, int(m.group(2)))}"', xml)


def _xml(obj):
    from openpyxl.xml.functions import tostring
    return tostring(obj.to_tree()).decode() if obj is not None else None


def _rules(ws):
    rules = []
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            rules.append((str(cf.sqref), rule.type, rule.operator, rule.priority, tuple(rule.formula or ()),
                          rule.stopIfTrue, rule.rank, rule.text, repr(rule.dxf), _xml(rule.colorScale),
                          _xml(rule.dataBar), _xml(rule.iconSet)))
    return sorted(rules, key=repr)


def _validations(ws):
    return sorted((str(dv.sqref), dv.type, dv.operator, dv.formula1, dv.formula2, dv.allow_blank,
                   dv.showErrorMessage, dv.showInputMessage, dv.error, dv.errorTitle, dv.prompt, dv.promptTitle,
                   dv.errorStyle) for dv in ws.data_validations.dataValidation)


def _sheet(ws):
    wb = ws.parent
    return {
        "autoFilterXml": _resolve_dxfs(wb, _xml(ws.auto_filter)) if ws.auto_filter.ref else None,
        "tableXml": sorted(_resolve_dxfs(wb, _xml(table)) for table in ws.tables.values()),
        "rules": _rules(ws),
        "validationDetails": _validations(ws),
        "columnWidths": sorted((k, v.width if v.customWidth else None) for k, v in ws.column_dimensions.items()
                               if v.customWidth),
        "rowHeights": sorted((k, v.height) for k, v in ws.row_dimensions.items() if v.height),
        "names": sorted((name, dn.attr_text) for name, dn in ws.defined_names.items()),
        "merged": sorted(map(str, ws.merged_cells.ranges)),
        "conditionalFormatting": sum(len(cf.rules) for cf in ws.conditional_formatting),
        "dataValidations": len(ws.data_validations.dataValidation),
        "tables": sorted(ws.tables),
        "charts": len(ws._charts),
        "images": len(ws._images),
        "pivots": len(ws._pivots),
        "freeze": ws.freeze_panes,
        "filter": ws.auto_filter.ref,
        "state": ws.sheet_state,
        "protection": repr(ws.protection),
        "hiddenColumns": sorted(k for k, v in ws.column_dimensions.items() if v.hidden),
        "hiddenRows": sorted(k for k, v in ws.row_dimensions.items() if v.hidden),
        "outlineRows": sorted(k for k, v in ws.row_dimensions.items() if v.outline_level),
        "printArea": ws.print_area,
        "printTitles": (ws.print_title_rows, ws.print_title_cols),
        "headerFooter": repr(ws.HeaderFooter.to_tree() is not None and
                             [(child.tag, child.text) for child in ws.HeaderFooter.to_tree()]),
        "breaks": ([b.id for b in ws.row_breaks.brk], [b.id for b in ws.col_breaks.brk]),
        "margins": repr(ws.page_margins),
        "pageSetup": (ws.page_setup.orientation, ws.page_setup.fitToWidth, ws.page_setup.paperSize),
        "printOptions": repr(ws.print_options),
        "view": (ws.sheet_view.zoomScale, ws.sheet_view.showGridLines, ws.sheet_view.rightToLeft),
        "tabColor": repr(ws.sheet_properties.tabColor),
    }


def _workbook(wb):
    return {
        "sheets": wb.sheetnames,
        "chartsheets": [cs.title for cs in wb.chartsheets],
        "names": sorted((name, dn.attr_text, dn.hidden) for name, dn in wb.defined_names.items()),
        "namedStyles": sorted(wb.named_styles),
        "properties": {name: getattr(wb.properties, name) for name in
                       ("title", "subject", "creator", "keywords", "category", "description")},
        "customProperties": [(p.name, str(p.value)) for p in wb.custom_doc_props],
        # An empty <workbookProtection/> (written by openpyxl) means no protection.
        "protection": {name: value for name, value in vars(wb.security).items() if value} if wb.security else {},
    }


@pytest.mark.parametrize("path", EXAMPLES, ids=lambda p: p.stem)
def test_example_fidelity(path, tmp_path):
    restored = tmp_path / path.name
    import_model(export_model(path), restored)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl's notes about extensions it drops itself
        original, copy = load_workbook(path, rich_text=True), load_workbook(restored, rich_text=True)

    assert _workbook(copy) == _workbook(original)
    for ws in original.worksheets:
        other = copy[ws.title]
        assert _sheet(other) == _sheet(ws), ws.title
        for row in ws.iter_rows():
            for cell in row:
                assert _cell(other[cell.coordinate]) == _cell(cell), f"{ws.title}!{cell.coordinate}"


@pytest.mark.parametrize("path", EXAMPLES, ids=lambda p: p.stem)
def test_example_tables_are_valid(path):
    """Excel repairs a file whose table column names differ from the header
    cells, so the examples must get them right."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(path)
    for ws in wb.worksheets:
        for table in ws.tables.values():
            if not table.headerRowCount:
                continue
            header = next(ws.iter_rows(min_row=ws[table.ref][0][0].row, max_row=ws[table.ref][0][0].row,
                                       min_col=ws[table.ref][0][0].column,
                                       max_col=ws[table.ref][0][-1].column, values_only=True))
            assert [c.name for c in table.tableColumns] == [str(v) for v in header], (ws.title, table.name)
