"""Printer settings and warnings about parts that are not exported."""

import posixpath
import re
import zipfile
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from PIL import Image as PILImage

from xlsx2txt import diff_models, export_model, export_xlsx, import_dir
from xlsx2txt.converter import roundtrip_diff
from xlsx2txt.importer import _vba_archive
from xlsx2txt.package import unsupported_warnings

PRINTER_A = b"\x00printer-A" * 50
PRINTER_B = b"\x00printer-B" * 50


def _rewrite(path, changes, extra=None):
    """Rewrite parts of a zip: {name: function(text) -> text} plus extra parts."""
    with zipfile.ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    for name, change in changes.items():
        parts[name] = change(parts.get(name, b"").decode("utf-8")).encode("utf-8")
    parts.update(extra or {})
    with zipfile.ZipFile(path, "w") as target:
        for name, content in parts.items():
            target.writestr(name, content)


def _add_printer_settings_like_excel(path, sheet_number, content):
    """Attach printer settings the way Excel does (Default bin content type)."""
    sheet = f"xl/worksheets/sheet{sheet_number}.xml"
    rels = f"xl/worksheets/_rels/sheet{sheet_number}.xml.rels"
    part = f"xl/printerSettings/printerSettings{sheet_number}.bin"

    def sheet_xml(text):
        text = text.replace("<worksheet ", '<worksheet xmlns:r="http://schemas.openxmlformats.org/'
                            'officeDocument/2006/relationships" ', 1)
        return re.sub(r"(<pageMargins[^>]*/>)", r'\1<pageSetup paperSize="9" orientation="landscape" r:id="rId9"/>',
                      text, count=1)

    def rels_xml(text):
        rel = ('<Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
               f'relationships/printerSettings" Target="../printerSettings/printerSettings{sheet_number}.bin"/>')
        if not text:
            text = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    '</Relationships>')
        return text.replace("</Relationships>", rel + "</Relationships>")

    def content_types(text):
        if 'Extension="bin"' in text:
            return text
        return text.replace("<Default ", '<Default Extension="bin" ContentType="application/vnd.'
                            'openxmlformats-officedocument.spreadsheetml.printerSettings"/><Default ', 1)

    _rewrite(path, {sheet: sheet_xml, rels: rels_xml, "[Content_Types].xml": content_types}, {part: content})


@pytest.fixture
def printer_xlsx(tmp_path):
    path = tmp_path / "printer.xlsx"
    wb = Workbook()
    wb.active.title = "First"
    wb.active["A1"] = "link"
    wb.active["A1"].hyperlink = "https://example.com"  # sheet already has relationships
    wb.create_sheet("Second")["A1"] = 2
    wb.create_sheet("Third")["A1"] = 3
    wb.save(path)
    _add_printer_settings_like_excel(path, 1, PRINTER_A)
    _add_printer_settings_like_excel(path, 2, PRINTER_B)
    _add_printer_settings_like_excel(path, 3, PRINTER_A)
    return path


def test_printer_settings_exported(printer_xlsx):
    model = export_model(printer_xlsx)
    names = [sheet.get("printerSettings") for sheet in model["sheets"]]
    assert names[0] == names[2] != names[1]  # identical settings are stored once
    assert sorted(model["printerSettings"].values()) == sorted([PRINTER_A, PRINTER_B])
    assert model["manifest"]["warnings"] == []


def test_printer_settings_restored(printer_xlsx, tmp_path):
    out_dir = tmp_path / "out"
    model = export_xlsx(printer_xlsx, out_dir)
    assert len(list((out_dir / "data" / "printer").iterdir())) == 2
    assert roundtrip_diff(model) == []

    restored = tmp_path / "restored.xlsx"
    import_dir(out_dir, restored)
    with zipfile.ZipFile(restored) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml").decode()
        rels = archive.read("xl/worksheets/_rels/sheet1.xml.rels").decode()
        types = archive.read("[Content_Types].xml").decode()
        rel_id = re.search(r'<pageSetup[^>]*r:id="([^"]+)"', sheet).group(1)
        target = re.search(rf'Id="{rel_id}"[^>]*Target="([^"]+)"', rels).group(1)
        assert archive.read(posixpath.normpath("xl/worksheets/" + target)) == PRINTER_A
        assert len(re.findall(r'Id="([^"]+)"', rels)) == len(set(re.findall(r'Id="([^"]+)"', rels)))
        assert types.count("spreadsheetml.printerSettings") == 3
    wb = load_workbook(restored)
    assert wb["First"]["A1"].hyperlink.target == "https://example.com"
    assert wb["First"].page_setup.orientation == "landscape"
    assert diff_models(model, export_model(restored), ignore_cached=True) == []


def test_printer_settings_with_macros(tmp_path):
    source = tmp_path / "macro.xlsm"
    wb = Workbook()
    wb.vba_archive = _vba_archive({"xl/vbaProject.bin": b"dummy"})
    wb.save(source)
    _add_printer_settings_like_excel(source, 1, PRINTER_A)

    out_dir = tmp_path / "out"
    model = export_xlsx(source, out_dir)
    assert list(model["printerSettings"].values()) == [PRINTER_A]
    restored = tmp_path / "restored.xlsm"
    import_dir(out_dir, restored)
    again = export_model(restored)
    assert again["vba"] == model["vba"]
    assert list(again["printerSettings"].values()) == [PRINTER_A]


def _png():
    buffer = BytesIO()
    PILImage.new("RGB", (4, 4), "red").save(buffer, format="PNG")
    return buffer.getvalue()


def test_warnings_for_lost_parts(tmp_path):
    path = tmp_path / "lost.xlsx"
    wb = Workbook()
    wb.active.title = "Data"
    wb.active.add_image(Image(BytesIO(_png())), "B2")
    wb.save(path)

    # A slicer list refers to slicer parts through relationships: not exported.
    slicer_list = ('<extLst><ext uri="{A8765BA9-456A-4dab-B4F3-ACF838C121DE}" '
                   'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
                   '<x14:slicerList><x14:slicer xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
                   'relationships" r:id="rId8"/></x14:slicerList></ext></extLst></worksheet>')
    # openpyxl writes drawings in the default namespace (Excel uses the xdr: prefix).
    # A shape linking to a relationship the drawing does not have cannot be kept.
    shape = ('<twoCellAnchor><from><col>5</col><colOff>0</colOff><row>5</row><rowOff>0</rowOff></from>'
             '<to><col>7</col><colOff>0</colOff><row>7</row><rowOff>0</rowOff></to>'
             '<sp><nvSpPr><cNvPr id="9" name="Arrow"><a:hlinkClick xmlns:a="http://schemas.openxmlformats.org/'
             'drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
             'r:id="rId9"/></cNvPr><cNvSpPr/></nvSpPr><spPr/></sp><clientData/>'
             '</twoCellAnchor></wsDr>')
    _rewrite(path, {
        "xl/worksheets/sheet1.xml": lambda text: text.replace("</worksheet>", slicer_list),
        "xl/drawings/drawing1.xml": lambda text: text.replace("</wsDr>", shape),
    }, {
        "xl/slicers/slicer1.xml": b"<slicers/>",
        "xl/connections.xml": b"<connections/>",
        "xl/unknownThing/part1.xml": b"<x/>",
    })

    warnings = unsupported_warnings(path)
    assert "slicers: 1 part(s) are not exported" in warnings
    assert "data connections / queries: 1 part(s) are not exported" in warnings
    assert "1 unknown part(s) are not exported: xl/unknownThing/part1.xml" in warnings
    assert "Sheet 'Data': slicers are not exported" in warnings
    assert "Sheet 'Data': 1 form control(s) or shape(s) with unsupported links are not exported" in warnings

    model = export_model(path)
    assert model["manifest"]["warnings"][:len(warnings)] == warnings
    # openpyxl's "extension is not supported" notes are replaced by the warnings above.
    assert not any("extension is not supported" in w for w in model["manifest"]["warnings"])
    assert len(model["media"]) == 1  # the picture next to the shape is still exported


def test_no_warnings_for_supported_files():
    examples = sorted((__import__("pathlib").Path(__file__).parent.parent / "examples").glob("*.xlsx"))
    assert examples
    for path in examples:
        assert unsupported_warnings(path) == [], path.name
