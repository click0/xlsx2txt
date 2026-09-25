"""Tests for images, charts, rich text and VBA sources."""

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from PIL import Image as PILImage

from xlsx2txt import diff_models, export_model, export_xlsx, import_dir, load_model
from xlsx2txt.compare import validate_model
from xlsx2txt.converter import roundtrip_diff
from xlsx2txt.drawings import anchor_from_json, anchor_to_json
from xlsx2txt.importer import _vba_archive


def _png(color, size=(20, 10)):
    buffer = BytesIO()
    PILImage.new("RGB", size, color).save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.fixture
def media_xlsx(tmp_path):
    path = tmp_path / "media.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for i in range(1, 6):
        ws.append([i, i * i])
    ws.add_image(Image(BytesIO(_png("red"))), "D2")

    bar = BarChart()
    bar.title = "Squares"
    bar.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    ws.add_chart(bar, "F2")

    ws["A8"] = CellRichText(["Plain ", TextBlock(InlineFont(b=True, color="FFFF0000"), "bold red"), " tail"])

    other = wb.create_sheet("Other")
    image = Image(BytesIO(_png("blue", (30, 30))))
    image.anchor = TwoCellAnchor(_from=AnchorMarker(col=1, row=1), to=AnchorMarker(col=3, row=4), editAs="oneCell")
    other.add_image(image)
    other.add_image(Image(BytesIO(_png("red"))), "H1")  # same content as on Data
    line = LineChart()
    line.add_data(Reference(ws, min_col=1, min_row=1, max_row=5))
    other.add_chart(line, "B10")
    wb.save(path)
    return path


def test_images_charts_rich_text_exported(media_xlsx):
    model = export_model(media_xlsx)
    data, other = model["sheets"]

    assert len(model["media"]) == 2  # identical images are stored once
    assert data["images"][0]["anchor"]["from"] == {"cell": "D2"}
    assert data["images"][0]["file"] in model["media"]
    # Drawings keep their anchors grouped by type: one-cell anchors come first.
    one_cell, two_cell = other["images"]
    assert one_cell["file"] == data["images"][0]["file"]
    assert two_cell["anchor"] == {
        "type": "twoCell", "from": {"cell": "B2"}, "to": {"cell": "D5"}, "editAs": "oneCell",
    }

    assert data["charts"][0]["anchor"]["from"] == {"cell": "F2"}
    assert "Squares" in data["charts"][0]["xml"]
    assert "lineChart" in other["charts"][0]["xml"]

    assert data["cells"]["A8"] == {
        "t": "rs",
        "v": ["Plain ", {"text": "bold red", "font": {"b": True, "color": "FFFF0000"}}, " tail"],
    }
    assert not model["manifest"]["warnings"]


def test_media_roundtrip(media_xlsx, tmp_path):
    out_dir = tmp_path / "media"
    model = export_xlsx(media_xlsx, out_dir)
    assert len(list((out_dir / "data" / "media").iterdir())) == 2
    assert roundtrip_diff(load_model(out_dir)) == []

    restored = tmp_path / "restored.xlsx"
    import_dir(out_dir, restored)
    assert diff_models(model, export_model(restored), ignore_cached=True) == []

    wb = load_workbook(restored, rich_text=True)
    assert len(wb["Data"]._charts) == 1 and len(wb["Other"]._charts) == 1
    assert len(wb["Data"]._images) == 1 and len(wb["Other"]._images) == 2
    assert str(wb["Data"]["A8"].value) == "Plain bold red tail"


def test_changed_image_shows_in_diff(media_xlsx):
    a = export_model(media_xlsx)
    b = export_model(media_xlsx)
    old = b["sheets"][0]["images"][0]["file"]
    b["media"]["image_new.png"] = _png("green")
    b["sheets"][0]["images"][0]["file"] = "image_new.png"
    changes = diff_models(a, b)
    assert "media/image_new.png: added" in changes
    assert any(line.startswith("[Data] images") and old in line for line in changes)


def test_missing_image_file_is_invalid(media_xlsx):
    model = export_model(media_xlsx)
    model["media"] = {}
    assert any("image file not found" in e for e in validate_model(model))


@pytest.mark.parametrize("anchor", [
    {"type": "oneCell", "from": {"cell": "C3", "offset": [10, 20]}, "size": [100, 200]},
    {"type": "twoCell", "from": {"cell": "A1"}, "to": {"cell": "E9", "offset": [5, 0]}},
    {"type": "absolute", "position": [1000, 2000], "size": [300, 400]},
])
def test_anchor_conversion(anchor):
    assert anchor_to_json(anchor_from_json(anchor)) == anchor


def test_vba_sources(tmp_path, monkeypatch):
    class FakeParser:
        def __init__(self, path):
            pass

        def detect_vba_macros(self):
            return True

        def extract_macros(self):
            yield "f", "VBA/Module1", "Module1.bas", 'Sub Hello()\r\n  MsgBox "hi"\r\nEnd Sub\r\n'
            yield "f", "VBA/ThisWorkbook", "ThisWorkbook.cls", "' workbook\r\n"

        def close(self):
            pass

    olevba = pytest.importorskip("oletools.olevba")
    monkeypatch.setattr(olevba, "VBA_Parser", FakeParser)

    source = tmp_path / "macro.xlsm"
    wb = Workbook()
    wb.vba_archive = _vba_archive({"xl/vbaProject.bin": b"dummy"})
    wb.save(source)

    out_dir = tmp_path / "macro"
    model = export_xlsx(source, out_dir)
    assert model["vbaSources"]["Module1.bas"] == 'Sub Hello()\n  MsgBox "hi"\nEnd Sub\n'
    assert (out_dir / "vba" / "modules" / "ThisWorkbook.cls").read_text() == "' workbook\n"

    loaded = load_model(out_dir)
    assert loaded["vbaSources"] == model["vbaSources"]
    assert loaded["vba"] == {"xl/vbaProject.bin": b"dummy"}  # sources are not packed into the file
