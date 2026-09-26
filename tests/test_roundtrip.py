"""Round-trip tests: xlsx -> model/directory -> xlsx."""

import datetime

import pytest
from openpyxl import Workbook, load_workbook

from xlsx2txt import build_workbook, diff_models, export_model, export_xlsx, import_dir, load_model
from xlsx2txt.converter import roundtrip_diff
from xlsx2txt.importer import _vba_archive
from xlsx2txt.storage import read_model

from tests.conftest import FIXTURES_DIR


def _reexport(model, tmp_path, suffix=".xlsx"):
    target = tmp_path / f"again{suffix}"
    build_workbook(model).save(target)
    return target, export_model(target, cached_values=False)


@pytest.mark.parametrize("name", ["simple.xlsx", "styled.xlsx", "empty.xlsx", "formula.xlsx", "complex.xlsx"])
def test_fixture_roundtrip(name):
    model = export_model(FIXTURES_DIR / name)
    assert roundtrip_diff(model) == []


def test_rich_roundtrip_via_directory(rich_xlsx, tmp_path):
    out_dir = tmp_path / "rich"
    original = export_xlsx(rich_xlsx, out_dir)
    restored = tmp_path / "restored.xlsx"
    import_dir(out_dir, restored)

    again = export_model(restored)
    assert diff_models(original, again, ignore_cached=True) == []
    assert read_model(out_dir)["sheets"] == original["sheets"]


def test_rich_values_and_types(rich_xlsx):
    model = export_model(rich_xlsx)
    cells = model["sheets"][0]["cells"]
    assert cells["A1"] == {"v": "Text", "t": "s"}
    assert cells["B1"] == {"v": 42, "t": "n"}
    assert cells["D1"] == {"v": True, "t": "b"}
    assert cells["E1"]["t"] == "d" and cells["E1"]["v"] == "2025-03-14T15:09:26"
    assert cells["G1"]["v"] == "12:30:00"
    assert cells["H1"] == {"v": "#N/A", "t": "e"}
    assert cells["A2"] == {"v": "=not a formula", "t": "s"}
    assert cells["B2"] == {"t": "f", "f": "=B1*2"}
    assert cells["C2"]["fType"] == "array" and cells["C2"]["fRef"] == "C2:C3"
    assert cells["A3"]["link"] == {"target": "https://example.com"}
    assert cells["A3"]["comment"] == {"text": "Check this", "author": "Reviewer"}
    assert "v" not in cells["C4"] and cells["C4"]["s"] > 0


def test_restored_workbook_content(rich_xlsx, tmp_path):
    path, _ = _reexport(export_model(rich_xlsx), tmp_path)
    wb = load_workbook(path)
    ws = wb["Main"]
    assert ws["A2"].value == "=not a formula" and ws["A2"].data_type == "s"
    assert ws["B2"].value == "=B1*2"
    assert ws["E1"].value == datetime.datetime(2025, 3, 14, 15, 9, 26)
    assert ws["B4"].font.name == "Arial" and ws["B4"].font.bold
    assert ws["B4"].font.color.theme == 4
    assert ws["B4"].number_format == "0.00%"
    assert ws["B4"].protection.locked is False
    assert "A6:C7" in [str(r) for r in ws.merged_cells.ranges]
    assert ws.freeze_panes == "B2"
    assert ws.column_dimensions["A"].width == 20
    assert ws.column_dimensions["B"].hidden
    assert ws.row_dimensions[2].height == 30
    assert wb["Hidden"].sheet_state == "hidden"
    assert wb.active.title == "Data <1>"
    assert "Items" in wb["Data <1>"].tables
    assert wb["Data <1>"].protection.sheet
    assert wb.defined_names["Total"].attr_text == "Main!$B$1"
    assert wb.properties.creator == "Tester"


def test_cached_values_exported(tmp_path):
    # openpyxl does not calculate formulas, so fake a cached value in the XML.
    import zipfile

    source = tmp_path / "calc.xlsx"
    wb = Workbook()
    wb.active["A1"] = 2
    wb.active["A2"] = "=A1*2"
    wb.save(source)
    patched = tmp_path / "patched.xlsx"
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(patched, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<f>A1*2</f><v />", b"<f>A1*2</f><v>4</v>")
            zout.writestr(item, data)

    model = export_model(patched)
    assert model["sheets"][0]["cells"]["A2"] == {"t": "f", "f": "=A1*2", "v": 4}
    assert export_model(patched, cached_values=False)["sheets"][0]["cells"]["A2"] == {"t": "f", "f": "=A1*2"}


def test_epoch_1904(tmp_path):
    from openpyxl.utils.datetime import CALENDAR_MAC_1904

    source = tmp_path / "mac.xlsx"
    wb = Workbook()
    wb.epoch = CALENDAR_MAC_1904
    wb.active["A1"] = datetime.datetime(2020, 5, 1)
    wb.save(source)

    model = export_model(source)
    assert model["workbook"]["epoch"] == 1904
    path, again = _reexport(model, tmp_path)
    assert diff_models(model, again, ignore_cached=True) == []
    assert load_workbook(path)["Sheet"]["A1"].value == datetime.datetime(2020, 5, 1)


def test_vba_roundtrip(tmp_path):
    source = tmp_path / "macro.xlsm"
    wb = Workbook()
    wb.active["A1"] = "macro"
    wb.vba_archive = _vba_archive({"xl/vbaProject.bin": b"\x00dummy-vba\xff"})
    wb.save(source)

    out_dir = tmp_path / "macro"
    model = export_xlsx(source, out_dir)
    assert model["vba"] == {"xl/vbaProject.bin": b"\x00dummy-vba\xff"}
    assert (out_dir / "vba" / "xl" / "vbaProject.bin").exists()

    restored = tmp_path / "restored.xlsm"
    import_dir(out_dir, restored)
    assert export_model(restored)["vba"] == model["vba"]


def test_load_model_accepts_file_and_dir(tmp_path):
    out_dir = tmp_path / "simple"
    export_xlsx(FIXTURES_DIR / "simple.xlsx", out_dir)
    assert diff_models(load_model(FIXTURES_DIR / "simple.xlsx"), load_model(out_dir)) == []


def test_custom_properties_roundtrip(tmp_path):
    import datetime as dt
    from openpyxl.packaging.custom import BoolProperty, DateTimeProperty, IntProperty, StringProperty

    source = tmp_path / "props.xlsx"
    wb = Workbook()
    for prop in (StringProperty(name="Project", value="Q1"), IntProperty(name="Build", value=7),
                 BoolProperty(name="Final", value=True),
                 DateTimeProperty(name="Due", value=dt.datetime(2026, 1, 2, 3, 4, 5))):
        wb.custom_doc_props.append(prop)
    wb.save(source)

    model = export_model(source)
    assert model["workbook"]["customProperties"][:2] == [
        {"name": "Project", "type": "StringProperty", "value": "Q1"},
        {"name": "Build", "type": "IntProperty", "value": 7},
    ]
    _, again = _reexport(model, tmp_path)
    assert again["workbook"]["customProperties"] == model["workbook"]["customProperties"]


def test_hidden_zero_width_column_roundtrip(tmp_path):
    import zipfile

    source = tmp_path / "cols.xlsx"
    wb = Workbook()
    wb.active["A1"] = 1
    wb.active.column_dimensions["B"].hidden = True
    wb.save(source)
    # Excel writes width="0" for such columns; openpyxl cannot, so patch the XML.
    patched = tmp_path / "patched.xlsx"
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(patched, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<col hidden="1" width="13"', b'<col hidden="1" width="0"')
            zout.writestr(item, data)

    model = export_model(patched)
    assert model["sheets"][0]["dimensions"]["columns"]["B"] == {"width": 0.0, "hidden": True}
    assert roundtrip_diff(model) == []


def test_theme_line_endings_preserved(tmp_path):
    out_dir = tmp_path / "out"
    model = export_xlsx(FIXTURES_DIR / "simple.xlsx", out_dir)
    model["theme"] = model["theme"].replace("\n", "\r\n")
    from xlsx2txt.storage import read_model, write_model

    write_model(model, out_dir)
    assert read_model(out_dir)["theme"] == model["theme"]


def test_zero_outline_levels_are_defaults(tmp_path):
    """LibreOffice writes outlineLevelCol="0"; openpyxl drops it on save."""
    import zipfile

    path = tmp_path / "outline.xlsx"
    wb = Workbook()
    wb.active["A1"] = 1
    wb.save(path)
    with zipfile.ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    parts["xl/worksheets/sheet1.xml"] = sheet.replace(
        "<sheetFormatPr ", '<sheetFormatPr outlineLevelRow="0" outlineLevelCol="0" ', 1).encode()
    with zipfile.ZipFile(path, "w") as target:
        for name, content in parts.items():
            target.writestr(name, content)
    model = export_model(path)
    assert "outlineLevelCol" not in model["sheets"][0].get("format", {})
    assert roundtrip_diff(model) == []
