"""Pivot tables: definitions and caches survive export/import."""

import importlib.util
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.pivot.table import TableDefinition
from openpyxl.xml.functions import fromstring, tostring

from xlsx2txt import diff_models, export_model, export_xlsx, import_dir
from xlsx2txt.compare import validate_model
from xlsx2txt.converter import roundtrip_diff
from xlsx2txt.text import render_text

EXAMPLES_DIR = Path(__file__).resolve().parent.parent / "examples"


def _generator():
    spec = importlib.util.spec_from_file_location("examples_generate", EXAMPLES_DIR / "generate.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture
def pivot_xlsx(tmp_path):
    path = tmp_path / "pivot.xlsx"
    _generator().pivot_table().save(path)
    return path


@pytest.fixture
def two_pivots_xlsx(tmp_path):
    """Two pivot tables on different sheets sharing one cache."""
    wb = _generator().pivot_table()
    first = wb["Data"]._pivots[0]
    second = TableDefinition.from_tree(fromstring(tostring(first.to_tree())))
    second.name = "Copy"
    second.cache = first.cache
    report = wb.create_sheet("Report")
    report.add_pivot(second)
    path = tmp_path / "two.xlsx"
    wb.save(path)
    return path


def test_pivot_exported(pivot_xlsx):
    model = export_model(pivot_xlsx)
    assert model["sheets"][0]["pivotTables"] == [{"table": "pivotTable1.xml", "cache": "pivotCache1"}]
    assert sorted(model["pivots"]) == ["pivotCache1.xml", "pivotCache1_records.xml", "pivotTable1.xml"]
    assert 'name="Totals"' in model["pivots"]["pivotTable1.xml"]
    assert '<worksheetSource ref="A1:B5" sheet="Data"' in model["pivots"]["pivotCache1.xml"]
    assert model["pivots"]["pivotTable1.xml"].count("\n") > 10  # indented, one element per line
    assert model["manifest"]["warnings"] == []


def test_pivot_roundtrip(pivot_xlsx, tmp_path):
    out_dir = tmp_path / "out"
    model = export_xlsx(pivot_xlsx, out_dir)
    assert (out_dir / "data" / "pivots" / "pivotTable1.xml").is_file()
    assert roundtrip_diff(model) == []

    restored = tmp_path / "restored.xlsx"
    import_dir(out_dir, restored)
    pivot = load_workbook(restored)["Data"]._pivots[0]
    assert pivot.name == "Totals"
    assert pivot.location.ref == "D1:E4"
    assert len(pivot.cache.records.r) == 4
    assert diff_models(model, export_model(restored), ignore_cached=True) == []


def test_shared_cache_stored_once(two_pivots_xlsx, tmp_path):
    model = export_model(two_pivots_xlsx)
    assert [s["pivotTables"][0]["cache"] for s in model["sheets"]] == ["pivotCache1", "pivotCache1"]
    assert sorted(model["pivots"]) == [
        "pivotCache1.xml", "pivotCache1_records.xml", "pivotTable1.xml", "pivotTable2.xml",
    ]
    assert roundtrip_diff(model) == []

    restored = tmp_path / "restored.xlsx"
    export_xlsx(two_pivots_xlsx, tmp_path / "out")
    import_dir(tmp_path / "out", restored)
    wb = load_workbook(restored)
    assert [ws._pivots[0].name for ws in wb.worksheets] == ["Totals", "Copy"]
    again = export_model(restored)
    assert sorted(again["pivots"]) == sorted(model["pivots"])  # still one shared cache


def test_missing_pivot_file_is_invalid(pivot_xlsx):
    model = export_model(pivot_xlsx)
    del model["pivots"]["pivotCache1.xml"]
    assert any("pivot file not found: data/pivots/pivotCache1.xml" in e for e in validate_model(model))


def test_changed_pivot_in_diff_and_text(pivot_xlsx):
    a = export_model(pivot_xlsx)
    b = export_model(pivot_xlsx)
    b["pivots"]["pivotTable1.xml"] = b["pivots"]["pivotTable1.xml"].replace('name="Totals"', 'name="Sums"')
    assert "pivots/pivotTable1.xml: changed" in diff_models(a, b)
    assert "pivot table Totals at D1:E4" in render_text(a)
