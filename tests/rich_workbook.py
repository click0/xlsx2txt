"""Builder of a feature-rich workbook used by round-trip tests."""

import datetime

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, GradientFill, PatternFill, Protection, Side
from openpyxl.styles.colors import Color
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo


def build_rich_workbook(path):
    wb = Workbook()
    wb.properties.creator = "Tester"
    wb.properties.title = "Rich workbook"

    ws = wb.active
    ws.title = "Main"
    ws["A1"] = "Text"
    ws["B1"] = 42
    ws["C1"] = 3.5
    ws["D1"] = True
    ws["E1"] = datetime.datetime(2025, 3, 14, 15, 9, 26)
    ws["E1"].number_format = "yyyy-mm-dd hh:mm:ss"
    ws["F1"] = datetime.date(2025, 1, 2)
    ws["F1"].number_format = "yyyy-mm-dd"
    ws["G1"] = datetime.time(12, 30)
    ws["G1"].number_format = "hh:mm"
    ws["H1"] = "#N/A"
    ws["I1"] = "Юнікод ✓"

    literal = ws["A2"]
    literal.value = "=not a formula"
    literal.data_type = "s"
    ws["B2"] = "=B1*2"
    ws["C2"] = ArrayFormula("C2:C3", "=B1:B2*2")

    ws["A3"] = "link"
    ws["A3"].hyperlink = "https://example.com"
    ws["A3"].comment = Comment("Check this", "Reviewer")

    styled = ws["B4"]
    styled.value = "styled"
    styled.font = Font(name="Arial", size=13, bold=True, italic=True, underline="single",
                       color=Color(theme=4, tint=-0.25))
    styled.fill = PatternFill("solid", fgColor="FFFFFF00")
    styled.border = Border(left=Side("thin", "FF000000"), bottom=Side("double", "FFFF0000"),
                           diagonal=Side("dashed"), diagonalUp=True)
    styled.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True,
                                 text_rotation=45, indent=1)
    styled.protection = Protection(locked=False, hidden=True)
    styled.number_format = "0.00%"

    ws["C4"].fill = GradientFill(stop=("FF0000FF", "FF00FF00"), degree=90)  # style only

    ws.merge_cells("A6:C7")
    ws["A6"] = "merged"
    ws["A6"].border = Border(top=Side("thick"))

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].hidden = True
    ws.column_dimensions["D"].font = Font(italic=True)
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[5].outline_level = 1
    ws.row_dimensions[5].hidden = True

    ws.freeze_panes = "B2"
    ws.sheet_properties.tabColor = "FF00AA00"
    ws.auto_filter.ref = "A1:I3"
    ws.print_area = "A1:I10"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"

    dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    dv.add("J1:J10")
    ws.add_data_validation(dv)

    ws.conditional_formatting.add(
        "B1:B10", CellIsRule(operator="greaterThan", formula=["10"], font=Font(color="FFFF0000"),
                             fill=PatternFill("solid", bgColor="FFFFC7CE")))
    ws.conditional_formatting.add(
        "C1:C10", ColorScaleRule(start_type="min", start_color="FFF8696B",
                                 end_type="max", end_color="FF63BE7B"))
    ws.conditional_formatting.add("D1:D10", FormulaRule(formula=["D1=TRUE"], stopIfTrue=True))

    data = wb.create_sheet("Data <1>")
    data.append(["Name", "Qty"])
    data.append(["a", 1])
    data.append(["b", 2])
    table = Table(displayName="Items", ref="A1:B3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    data.add_table(table)
    data.defined_names["LocalName"] = DefinedName("LocalName", attr_text="'Data <1>'!$B$2")
    data.protection.sheet = True
    data.protection.password = "secret"

    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "=Main!B1+1"

    wb.defined_names["Total"] = DefinedName("Total", attr_text="Main!$B$1")
    wb.active = 1
    wb.save(path)
    return path
